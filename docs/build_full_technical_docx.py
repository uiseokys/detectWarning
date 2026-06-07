from __future__ import annotations

import ast
import json
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "detectWarning_full_technical_documentation.docx"
PRESENTATION_VISUALS = DOCS / "presentation_visuals"
USER_VISUALS = DOCS / "_user_pdf_visuals_png"

FONT = "Malgun Gothic"
INK = "111827"
MUTED = "475569"
BLUE = "2563EB"
BLUE_DARK = "1E3A8A"
CYAN = "0891B2"
GREEN = "16A34A"
YELLOW = "D97706"
RED = "DC2626"
LIGHT_BLUE = "DBEAFE"
LIGHT_CYAN = "CFFAFE"
LIGHT_GREEN = "DCFCE7"
LIGHT_YELLOW = "FEF3C7"
LIGHT_RED = "FEE2E2"
LIGHT_GRAY = "F1F5F9"


KEY_RUNTIME_FILES = [
    "app/inference_server.py",
    "app/camera_uploader.py",
    "app/main.py",
    "app/action_realtime.py",
    "app/risk_analyzer.py",
    "app/stt_service.py",
    "app/webrtc_stream.py",
    "app/event_clip_service.py",
    "app/backend_bridge.py",
    "app/person_classifier.py",
    "app/detector.py",
    "app/tracker.py",
    "app/runtime_state.py",
    "app/system_monitoring.py",
    "app/app_backend_client.py",
    "app/remote_inference.py",
    "app/audio_detector.py",
    "app/event_logger.py",
]

KEY_TRAINING_FILES = [
    "app/action_training_pipeline.py",
    "app/guideline_pose_dataset.py",
    "app/extract_rgb_video_features.py",
    "app/action_model.py",
    "app/train_fight_bilstm.py",
    "app/specialized_action_tasks.py",
    "app/hybrid_pose_ensemble.py",
    "app/ensemble_action_models.py",
    "app/auto_tune_action_training.py",
    "app/fight_bilstm_seed_sweep.py",
    "app/gpu_autotune.py",
    "app/pipeline_prepare.py",
    "app/action_training_pipeline_mean_max.py",
    "app/mean_max_pooling.py",
    "app/predownload_aihub_filekey.py",
    "app/cleanup_transient_data.py",
]

KEY_DASHBOARD_FILES = [
    "app/training_dashboard.py",
    "app/training_dashboard_view.py",
    "app/dashboard_runtime.py",
    "app/dashboard_aihub.py",
    "app/dashboard_quality.py",
    "app/dashboard_normal_ratio.py",
    "app/dashboard_notifications.py",
    "app/dashboard_gpu.py",
    "app/dashboard_performance_plan.py",
    "app/training_insights.py",
    "app/reporting.py",
    "app/training_config.py",
    "app/update_pages_site.py",
    "app/run_training_share.py",
    "app/runtime_health_check.py",
]


FILE_NOTES = {
    "app/inference_server.py": [
        "실사용 추론 서버의 중심 파일이다. FastAPI 앱, 추론 대시보드, 카메라 클라이언트 세션, 프레임 분석 API, 오디오 분석 API, WebRTC offer 처리, 위험 이벤트 생성, 백엔드 전송, 위험 기록 초기화를 한 곳에서 조율한다.",
        "ClientSession 단위로 최신 원본 프레임, 분석 프레임, CCTV 코드, 클라이언트 이름, 위험 점수, 음성 상태, 백엔드 동기화 상태를 유지한다.",
        "camera_uploader.py가 /analyze/frame과 /analyze/audio로 보내는 데이터를 받고 detector, action_realtime, stt_service, risk_analyzer, event_clip_service, webrtc_stream, backend_bridge를 연결한다.",
        "백엔드 이벤트는 riskScore, riskClass, riskLevel, videoReason, audioRiskSignalDetected, clipUrl 같은 메타데이터만 전송한다. STT 원문과 matchedKeywords는 백엔드로 보내지 않는다.",
    ],
    "app/camera_uploader.py": [
        "Mac/Windows 카메라 또는 iPhone 연동 카메라를 추론 서버에 붙이는 경량 클라이언트다.",
        "같은 장치에서 실행하면 client_id를 보존해 CCTV 연결 코드가 매번 바뀌지 않게 한다.",
        "프레임은 서버가 요구하는 크기와 FPS에 맞춰 JPEG로 업로드하고, STT가 켜져 있을 때만 오디오 조각을 별도 업로드한다.",
        "카메라 장치 목록 탐색, 재연결, FPS 제한, 오디오 스트림 전송을 담당한다.",
    ],
    "app/main.py": [
        "서버 없이 로컬 데스크톱에서 웹캠 위험 감지를 바로 시연할 수 있는 실행 진입점이다.",
        "OpenCV 카메라 입력, YOLO 사람 검출, pose overlay, 위험 점수 표시, 음성 위험 감지를 한 화면에 표시한다.",
        "대시보드 서버형 운영 전 단계의 단독 시연 또는 디버깅용 흐름이다.",
    ],
    "app/action_realtime.py": [
        "실시간 행동 인식 계층이다. 학습된 pose/RGB/I3D 모델과 fight/fall BiLSTM 보조 모델을 읽어 RealtimeActionResult로 정리한다.",
        "기본 clip_seconds=4, min_interval=1.5초, pose sequence=32, RGB frame=16, image_size=112 기준으로 최근 프레임을 묶어 판단한다.",
        "violence는 높은 abnormal_score와 confidence를 요구하고, collapse는 정적 자세와 temporal abnormal을 함께 보며, loitering은 일정 시간 이상 반복 hit와 이동 패턴을 요구한다.",
        "이 파일의 출력은 risk_analyzer.py가 최종 위험 점수로 바꾸는 영상 기반 근거가 된다.",
    ],
    "app/risk_analyzer.py": [
        "영상 점수, 음성 점수, 결합 점수, 위험 레벨, 위험 근거를 계산하는 의사결정 계층이다.",
        "SUPPORTED_ACTION_LABELS는 violence, collapse, loitering이다. abduction은 현재 추론 핵심 클래스에서 제외되어야 한다.",
        "클래스별 ACTION_SCORE_POLICIES로 audio_bonus, min_confidence, min_abnormal, weak_signal_multiplier, video_only_cap을 분리해 오탐을 줄인다.",
        "videoOnlyScore와 최종 riskScore를 분리해 음성이 있을 때 얼마나 위험 점수가 올라갔는지 발표에서 설명할 수 있게 한다.",
    ],
    "app/stt_service.py": [
        "CLOVA Speech Recognition 호출, 오디오 전처리, 비용/지연 제어, 빈 응답 처리, 실패 fallback을 담당한다.",
        "duration, RMS, peak, active_ratio gate를 통과한 오디오만 CLOVA에 보낸다. 너무 짧거나 조용한 조각은 바로 스킵해 비용과 지연을 줄인다.",
        "raw, normalized, boosted 후보를 순차 시도해 영상 파일 음성이 작거나 압축된 경우의 인식률을 보완한다.",
        "STT 결과 원문은 risk_analyzer 내부 위험 판단에만 쓰고 백엔드에는 boolean/score 메타데이터만 전달한다.",
    ],
    "app/webrtc_stream.py": [
        "분석 대시보드에 보이는 최신 processed frame을 WebRTC video track으로 송출한다.",
        "H.264를 우선 codec으로 쓰고 VP8 fallback을 둔다. 최신 프레임이 잠시 늦어져도 stale frame hold로 화면 깜빡임을 줄인다.",
        "백엔드는 JPEG를 WebRTC로 변환하지 않고 offer만 중계하며, 추론 서버가 직접 answer를 만든다.",
    ],
    "app/event_clip_service.py": [
        "클라이언트별 분석 프레임을 최근 20~30초 링버퍼로 유지하고, 위험 이벤트 발생 시 전후 구간 MP4를 만든다.",
        "기본 의도는 판단 전 5초 + 판단 후 10초, 총 15초 분석 프레임 클립이다.",
        "iPhone Safari/PWA 재생을 위해 Range 요청을 지원하는 clip endpoint와 연결된다.",
    ],
    "app/backend_bridge.py": [
        "추론 서버가 백엔드 장애 때문에 느려지지 않도록 전송 회로 차단기를 제공한다.",
        "연속 실패가 누적되면 짧은 시간 circuit을 열어 과도한 재시도와 지연을 막고, 성공 시 실패 상태를 복구한다.",
    ],
    "app/person_classifier.py": [
        "검출된 bbox가 실제 사람인지 검증하고 화면 overlay 품질을 정리한다.",
        "사람 #1 같은 의미 없는 번호 라벨은 제거하고, 화면에는 직관적으로 '사람'만 표시하도록 조정된 흐름과 연결된다.",
    ],
    "app/detector.py": [
        "YOLO 기반 사람/pose 검출 래퍼다. confidence threshold, image size, detector device 같은 실시간 검출 성능 파라미터의 영향을 직접 받는다.",
        "검출 결과는 tracker, person_classifier, action_realtime으로 전달된다.",
    ],
    "app/tracker.py": [
        "프레임 사이 사람 bbox를 연결해 이동량, 머무름, 객체 수, 접촉 의심 같은 시간적 근거를 제공한다.",
        "loitering 오탐 완화에서 같은 위치에 정지한 사람을 배회로 보지 않기 위한 입력이 된다.",
    ],
    "app/runtime_state.py": [
        "ClientSession 상태를 API 응답용 안전한 JSON으로 변환한다.",
        "NaN/Inf, 비직렬화 객체, 오래된 프레임 상태 때문에 대시보드가 깨지는 것을 막는 방어 계층이다.",
    ],
    "app/system_monitoring.py": [
        "CPU, RAM, GPU, 서버 health, WebRTC 상태를 대시보드와 발표 모드에 표시할 수 있게 수집한다.",
        "추론 서버의 자원 여유를 확인하고 FPS/분석 설정을 올릴 근거가 된다.",
    ],
    "app/action_training_pipeline.py": [
        "AIHub 다운로드부터 split, pose prepare, 학습, 평가, 결과 저장까지 이어지는 학습 파이프라인 중심 파일이다.",
        "pipeline_status.json, manifests, artifacts, metrics, skip report를 만들고 대시보드가 이를 읽는다.",
        "학습 대시보드가 만든 runtime config를 받아 stage별로 실행된다.",
    ],
    "app/guideline_pose_dataset.py": [
        "AIHub XML guideline을 반영해 event 구간, 앞뒤 context, normal clip, filekey 추적 정보를 가진 guideline_prepared manifest를 만든다.",
        "pose 실패 clip도 RGB/I3D-only fallback으로 살릴 수 있게 manifest를 구성하는 핵심 전처리 계층이다.",
    ],
    "app/extract_rgb_video_features.py": [
        "guideline clip 단위로 RGB/I3D 또는 fallback RGB feature를 추출해 npz와 manifest 필드를 갱신한다.",
        "filekey, split, label, sample_weight, rgb_feature_path 추적성이 이후 학습/진단의 기준이 된다.",
    ],
    "app/action_model.py": [
        "pose sequence 분류 모델과 학습 루프를 담당한다. PoseSequenceDataset, TemporalPoseClassifier, MeanMaxTemporalPoseClassifier, FocalLoss가 핵심이다.",
        "balanced sampler, class weight, focal loss, validation metric, checkpoint 저장을 수행한다.",
    ],
    "app/train_fight_bilstm.py": [
        "짧은 fight/noFight 또는 fall/normal 데이터셋을 CNN feature + BiLSTM + attention 구조로 빠르게 미세조정한다.",
        "기존 3클래스 모델을 대체하는 주 모델이 아니라 violence/collapse 보조 점수 또는 조건부 booster로 쓰는 방향이 안정적이다.",
    ],
    "app/specialized_action_tasks.py": [
        "normal 포함 감지 task와 abnormal class 분류 task를 분리해 학습/평가할 수 있게 한다.",
        "감지 목적과 세부 분류 목적이 다를 때 threshold와 class bias를 따로 잡기 위한 계층이다.",
    ],
    "app/hybrid_pose_ensemble.py": [
        "pose, RGB/I3D, bbox/meta feature를 결합해 classical model 기반 hybrid 후처리 성능을 탐색한다.",
        "RGB/I3D-only row와 label 충돌 row를 필터링하고 feature matrix를 만들어 검증한다.",
    ],
    "app/ensemble_action_models.py": [
        "여러 trial checkpoint의 probability를 weighted average, bias, temperature 방식으로 앙상블한다.",
        "단일 모델 성능이 불안정할 때 macro F1 안정화를 목표로 한다.",
    ],
    "app/auto_tune_action_training.py": [
        "여러 학습 설정 trial을 반복 실행하고 가장 좋은 objective score를 고른다.",
        "class weight, focal loss, seed, normal ratio, fusion 설정을 바꿔가며 자동 탐색한다.",
    ],
    "app/training_dashboard.py": [
        "학습 대시보드의 FastAPI 서버다. 파일키 추천, 자동 추출, 병렬 다운로드, 재시작, 학습 시작, 로그, 결과 표시 API를 제공한다.",
        "파일이 매우 크므로 UI 렌더링은 training_dashboard_view.py, AIHub 추천은 dashboard_aihub.py, 런타임 보조는 dashboard_runtime.py로 일부 분리되어 있다.",
    ],
    "app/training_dashboard_view.py": [
        "학습 대시보드 HTML/CSS/JS 렌더링을 담당한다.",
        "분포 표, 로그, 그래프, 작업 제어, 학습 결과 시각화가 이 파일을 통해 사용자에게 표시된다.",
    ],
    "app/dashboard_aihub.py": [
        "AIHub filekey의 class, inside/outside, coverage, 추천 우선순위를 판단한다.",
        "현재 정책은 violence는 outside 우선, collapse/loitering은 outside+inside 활용, abduction은 부족하면 제외하거나 제한하는 방향과 연결된다.",
    ],
    "app/dashboard_runtime.py": [
        "대시보드 subprocess 실행, 로그 tail, job snapshot, Pages sync 같은 운영 보조 로직을 담당한다.",
    ],
    "app/training_insights.py": [
        "학습 결과를 사람이 이해하기 쉬운 진단 메시지로 바꾼다.",
        "accuracy, macro F1, validation loss, confusion matrix, class report를 보고 과적합/클래스 불균형/오탐 가능성을 설명한다.",
    ],
    "app/reporting.py": [
        "JSON atomic write, manifest 통계, metrics 읽기, confusion matrix 정리 등 결과 저장/조회 공통 계층이다.",
        "상태 파일이 읽는 중 깨지는 문제를 줄이기 위해 atomic write와 retry 흐름이 중요하다.",
    ],
}


def set_run_font(run, size: float | None = None, *, bold: bool = False, color: str = INK) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_text(cell, text: str, *, fill: str | None = None, bold: bool = False, color: str = INK, size: float = 8.0) -> None:
    cell.text = ""
    if fill:
        shade(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    set_run_font(run, size, bold=bold, color=color)


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(4)

    for style_name, size, color in [
        ("Heading 1", 15.5, BLUE),
        ("Heading 2", 12.2, BLUE_DARK),
        ("Heading 3", 10.5, CYAN),
    ]:
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("detectWarning 기술 상세 문서")
    set_run_font(run, 8, color=MUTED)


def add_paragraph(doc: Document, text: str, *, size: float = 9.4, bold: bool = False, color: str = INK) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, size, bold=bold, color=color)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(item)
        set_run_font(r, 9.0)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(item)
        set_run_font(r, 9.0)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], *, widths: list[float] | None = None, font_size: float = 7.6) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell_text(table.cell(0, idx), header, fill=LIGHT_BLUE, bold=True, color=BLUE_DARK, size=8.0)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell_text(cells[idx], value, size=font_size)
            if widths and idx < len(widths):
                cells[idx].width = Inches(widths[idx])
    doc.add_paragraph()


def add_callout(doc: Document, title: str, body: str, *, fill: str = LIGHT_YELLOW) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_run_font(r, 9.5, bold=True, color=BLUE_DARK)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_run_font(r2, 8.7)
    doc.add_paragraph()


def add_image(doc: Document, path: Path, caption: str, *, max_width: float = 7.05, max_height: float = 4.3) -> None:
    if not path.exists():
        return
    with Image.open(path) as image:
        width_px, height_px = image.size
    aspect = width_px / max(height_px, 1)
    width = max_width
    height = width / aspect
    if height > max_height:
        height = max_height
        width = height * aspect
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(path), width=Inches(width), height=Inches(height))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = c.add_run(caption)
    set_run_font(cr, 7.8, color=MUTED)


def add_landscape_section(doc: Document) -> None:
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(11)
    sec.page_height = Inches(8.5)
    sec.top_margin = Inches(0.4)
    sec.bottom_margin = Inches(0.4)
    sec.left_margin = Inches(0.4)
    sec.right_margin = Inches(0.4)


def add_portrait_section(doc: Document) -> None:
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.6)
    sec.left_margin = Inches(0.65)
    sec.right_margin = Inches(0.65)


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig", errors="replace")


def ast_summary(path: Path) -> dict:
    text = read_text(path)
    try:
        tree = ast.parse(text)
    except SyntaxError:
        return {"lines": text.count("\n") + 1, "classes": [], "functions": [], "imports": []}

    classes = [node.name for node in tree.body if isinstance(node, ast.ClassDef)]
    functions = [node.name for node in tree.body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))]
    imports: set[str] = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.ImportFrom):
            module = node.module or ""
            if module.startswith("app."):
                imports.add("app/" + module.split(".", 1)[1].replace(".", "/") + ".py")
            elif node.level > 0:
                for alias in node.names:
                    candidate = "app/" + alias.name.replace(".", "/") + ".py"
                    if (ROOT / candidate).exists():
                        imports.add(candidate)
        elif isinstance(node, ast.Import):
            for alias in node.names:
                if alias.name.startswith("app."):
                    imports.add("app/" + alias.name.split(".", 1)[1].replace(".", "/") + ".py")
    return {
        "lines": text.count("\n") + 1,
        "classes": classes,
        "functions": functions,
        "imports": sorted(imports),
    }


def file_detail_rows(files: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for file in files:
        path = ROOT / file
        if not path.exists():
            continue
        summary = ast_summary(path)
        notes = FILE_NOTES.get(file, ["보조 모듈이다.", "상위 파이프라인 또는 대시보드에서 호출된다.", "상태 파일, manifest, API 응답의 안정성을 보완한다."])
        rows.append([
            file,
            str(summary["lines"]),
            "\n".join(notes[:2]),
            "\n".join(notes[2:]),
            ", ".join(summary["classes"][:8]) or "-",
            ", ".join(summary["functions"][:10]) or "-",
        ])
    return rows


def dependency_rows(files: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for file in files:
        path = ROOT / file
        if not path.exists():
            continue
        summary = ast_summary(path)
        local_imports = [item for item in summary["imports"] if (ROOT / item).exists()]
        if local_imports:
            rows.append([file, "\n".join(local_imports[:14])])
    return rows


def all_app_rows() -> list[list[str]]:
    rows: list[list[str]] = []
    for path in sorted((ROOT / "app").glob("*.py")):
        rel = path.relative_to(ROOT).as_posix()
        summary = ast_summary(path)
        first_note = FILE_NOTES.get(rel, ["지원 모듈 또는 실행 스크립트다."])[0]
        rows.append([
            rel,
            str(summary["lines"]),
            first_note,
            ", ".join(summary["classes"][:6]) or "-",
            ", ".join(summary["functions"][:8]) or "-",
        ])
    return rows


def risk_rules_overview() -> list[list[str]]:
    path = ROOT / "configs" / "risk_rules.json"
    if not path.exists():
        return []
    data = json.loads(read_text(path))
    rows: list[list[str]] = []
    for rule in data.get("category_rules", []):
        rows.append([
            rule.get("id", "-"),
            rule.get("name", "-"),
            str(rule.get("base_score", "-")),
            ", ".join(rule.get("patterns", [])[:4]),
            rule.get("description", ""),
        ])
    return rows


def test_rows() -> list[list[str]]:
    rows: list[list[str]] = []
    for path in sorted((ROOT / "tests").glob("test_*.py")):
        rel = path.relative_to(ROOT).as_posix()
        summary = ast_summary(path)
        rows.append([
            rel,
            str(summary["lines"]),
            ", ".join(summary["functions"][:10]) or "-",
            "위험 판단, 백엔드 연동, WebRTC, 학습/대시보드 보조 로직이 깨지지 않는지 확인하는 회귀 테스트다.",
        ])
    return rows


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    r = p.add_run("detectWarning 기술 상세 문서")
    set_run_font(r, 26, bold=True, color=INK)
    p2 = doc.add_paragraph()
    r2 = p2.add_run("학습 파이프라인, 실시간 추론 서버, 음성 위험 감지, WebRTC, 백엔드/PWA 연동, 파일별 책임과 상호작용")
    set_run_font(r2, 13, bold=True, color=BLUE)
    add_table(
        doc,
        ["항목", "내용"],
        [
            ["작성일", datetime.now().strftime("%Y-%m-%d")],
            ["문서 목적", "프로젝트 코드의 각 파일이 어떤 역할을 맡고, 어떤 기준으로 동작하며, 다른 파일 및 로직과 어떻게 상호작용하는지 설명한다."],
            ["대상 범위", "AIHub 학습/전처리, pose/RGB/I3D/BiLSTM 학습, 자동 튜닝, 실시간 감지, CLOVA STT, 위험 점수 결합, WebRTC, 위험 클립, 앱 백엔드 연동"],
            ["현재 핵심 클래스", "violence, collapse, loitering. normal은 기준 상태이며, abduction은 현재 추론 서버 핵심 위험 클래스에서 제외된 상태로 정리한다."],
            ["개인정보 원칙", "STT 원문, transcript, matchedKeywords는 앱 백엔드로 전송하지 않고 위험 신호 여부와 점수 메타데이터만 전송한다."],
        ],
        widths=[1.6, 5.6],
        font_size=8.3,
    )


def build() -> None:
    doc = Document()
    setup_document(doc)
    add_title(doc)

    doc.add_heading("1. 전체 시스템 목적과 설계 원칙", level=1)
    add_paragraph(
        doc,
        "detectWarning은 CCTV, 웹캠, 업로드 영상에서 위험 행동과 위험 음성 신호를 함께 분석하는 실시간 위험상황 감지 시스템이다. "
        "학습 단계에서는 AIHub 이상행동 CCTV 데이터를 guideline XML 기준으로 event clip 단위로 정리하고, pose sequence와 RGB/I3D feature를 추출해 모델을 학습한다. "
        "운영 단계에서는 카메라 클라이언트가 추론 서버에 프레임과 오디오를 보내고, 추론 서버가 영상 모델과 CLOVA STT 기반 음성 위험 분석을 결합해 위험 점수와 근거를 만든다."
    )
    add_bullets(
        doc,
        [
            "기능 보존: 학습, 대시보드, 추론, 백엔드 연동이 이미 동작하므로 구조를 갈아엎기보다 파일별 역할을 보존한 상태에서 안정화했다.",
            "오탐 완화: 단일 프레임 또는 단일 키워드만으로 경고하지 않고 클래스별 threshold, 연속 hit, cooldown, hysteresis, state machine을 적용한다.",
            "음성 기여도 설명: videoOnlyScore와 최종 riskScore를 분리해 음성이 있을 때 위험 점수가 얼마나 상승했는지 보여준다.",
            "개인정보 보호: 위험 판단에 STT 원문은 쓰지만 백엔드/PWA에는 transcript를 보내지 않는다.",
            "운영 반응성: 영상은 WebRTC direct-first로 백엔드를 거치지 않게 하고, 상태와 이벤트는 push/SSE/API 메타데이터 중심으로 전달한다.",
        ]
    )
    add_callout(
        doc,
        "핵심 설계 한 줄 요약",
        "영상 모델이 위험 후보를 만들고, 음성 위험 신호가 그 후보를 강화하며, 최종 이벤트는 시간 안정화와 백엔드 전송 정책을 통과한 경우에만 기록/알림으로 확정된다.",
        fill=LIGHT_GREEN,
    )

    doc.add_heading("2. 시스템 전체 흐름", level=1)
    add_image(doc, USER_VISUALS / "detectwarning_architecture_ppt_white.png", "전체 아키텍처: 추론 서버, 카메라 클라이언트, 앱 백엔드, PWA가 연결되는 구조", max_width=7.05, max_height=4.2)
    add_image(doc, PRESENTATION_VISUALS / "06_demo_storyboard.png", "발표 시연 스토리보드: 카메라 입력부터 위험 이벤트 확인까지", max_width=7.05, max_height=4.0)
    add_numbered(
        doc,
        [
            "카메라 클라이언트 또는 영상 테스트 클라이언트가 추론 서버에 연결되고, 클라이언트별 고유 CCTV 코드가 생성 또는 재사용된다.",
            "프레임은 추론 서버의 ClientSession에 저장되고, detector/person_classifier/tracker/action_realtime을 거쳐 영상 위험 후보가 계산된다.",
            "오디오가 활성화되어 있으면 stt_service가 CLOVA CSR에 보낼 수 있는 품질인지 gate하고, 통과한 음성만 STT 처리한다.",
            "risk_analyzer가 영상 점수와 음성 점수를 결합해 videoOnlyScore, audioScore, audioVideoGain, riskScore, riskLevel, riskEvidence를 만든다.",
            "inference_server.py의 백엔드 전송 정책이 클래스별 threshold, confirm hit, window, cooldown을 확인한다.",
            "확정 이벤트는 event_clip_service가 생성한 clipUrl과 함께 앱 백엔드 /inference/events로 전송된다.",
            "PWA는 백엔드 DB/푸시 알림/SSE 상태와 추론 서버 WebRTC 분석 프레임을 조합해 사용자에게 위험상황을 보여준다.",
        ]
    )

    doc.add_heading("3. 실시간 추론 서버 파일별 상세 책임", level=1)
    add_table(
        doc,
        ["파일", "라인", "담당 로직", "상호작용/출력", "주요 클래스", "주요 함수"],
        file_detail_rows(KEY_RUNTIME_FILES),
        widths=[1.35, 0.45, 2.55, 2.55, 1.25, 1.8],
        font_size=7.1,
    )

    doc.add_heading("4. 영상 위험 행동 분석 기준", level=1)
    add_table(
        doc,
        ["항목", "현재 기준", "선택 이유"],
        [
            ["분석 대상 클래스", "violence, collapse, loitering", "현재 학습/추론에서 안정적으로 쓰는 3개 위험 클래스다. abduction은 학습 제외 상태이므로 추론 서버 위험 클래스에서 빼는 것이 혼동을 줄인다."],
            ["기본 clip 길이", "4초", "행동은 단일 프레임보다 전후 움직임이 중요하다. 4초는 실시간 지연을 과도하게 늘리지 않으면서 행동 흐름을 볼 수 있는 절충값이다."],
            ["최소 분석 간격", "1.5초", "매 프레임 추론하면 GPU/CPU 사용량과 오탐이 늘어난다. 1.5초 간격은 반응성과 안정성의 균형값이다."],
            ["pose sequence 길이", "32 frame", "pose 기반 행동 변화량을 학습 모델 입력 길이에 맞춰 안정적으로 제공한다."],
            ["RGB frame 수", "16 frame", "I3D/RGB 계열 모델이 짧은 영상 문맥을 보기에 충분하고 실시간 처리 부담을 낮춘다."],
            ["violence guard", "min_confidence 0.74, min_abnormal 0.86, weak_signal_multiplier 0.42", "폭력은 오탐 비용이 크므로 약한 영상 신호는 강하게 낮추고, 사람 수/음성/반복 hit가 있을 때만 강화한다."],
            ["collapse guard", "min_confidence 0.58, min_abnormal 0.76, temporal_abnormal 0.94", "쓰러짐은 confidence가 낮을 수 있어 threshold를 너무 높이지 않되, 시간적 자세 변화와 정적 상태를 같이 본다."],
            ["loitering guard", "min_duration 8초, min_repeated_hits 2, video_only_cap 62", "서 있거나 기다리는 사람을 배회로 오탐하지 않게 지속 시간과 반복 hit를 요구한다."],
        ],
        widths=[1.8, 2.4, 4.5],
        font_size=7.8,
    )

    doc.add_heading("5. 음성 위험 감지와 CLOVA STT 로직", level=1)
    add_image(doc, PRESENTATION_VISUALS / "05_audio_privacy_boundary.png", "음성 개인정보 경계: STT 원문은 추론 서버 내부에서만 사용", max_width=7.05, max_height=4.0)
    add_table(
        doc,
        ["단계", "기준", "담당 파일/함수", "이유"],
        [
            ["오디오 품질 gate", "duration >= 0.55초, RMS >= 0.006, peak >= 0.035, active_ratio >= 0.018", "stt_service.py", "너무 짧거나 작은 음성은 CLOVA 비용만 발생하고 빈 응답 가능성이 높아 사전에 거른다."],
            ["CLOVA 요청 후보", "raw, normalized RMS 0.075, boosted RMS 0.11", "clova_audio_candidates", "영상 파일/맥북 마이크처럼 음량이 낮거나 압축된 경우 인식률을 올리기 위한 단계적 보정이다."],
            ["요청 속도 제한", "최소 3초 간격, 분당 최대 18초 오디오", "should_accept_audio_upload", "CLOVA 비용과 지연을 제어하면서 실시간 위험 신호는 놓치지 않기 위한 절충값이다."],
            ["위험 키워드/문맥", "A1~A10 category_rules", "configs/risk_rules.json + risk_analyzer.py", "협박, 구조요청, 공포 반응, 통제/강압, 충돌음을 일반 대화와 분리해 점수화한다."],
            ["백엔드 전송 제한", "transcript/matchedKeywords 전송 금지", "inference_server.py payload builder", "앱에는 위험 신호 여부와 점수만 필요하므로 개인정보 노출을 줄인다."],
        ],
        widths=[1.6, 2.3, 2.0, 3.0],
        font_size=7.7,
    )
    rows = risk_rules_overview()
    if rows:
        add_table(doc, ["ID", "카테고리", "기본점수", "대표 패턴", "설명"], rows, widths=[0.6, 1.8, 0.8, 3.0, 2.8], font_size=6.9)

    doc.add_heading("6. 영상+음성 결합 점수와 이벤트 확정", level=1)
    add_image(doc, PRESENTATION_VISUALS / "01_audio_lift_score.png", "영상만 사용한 점수와 영상+음성 결합 점수 비교", max_width=7.05, max_height=4.0)
    add_image(doc, PRESENTATION_VISUALS / "03_fusion_matrix.png", "영상/음성 신호 조합에 따른 최종 위험 판단", max_width=7.05, max_height=4.0)
    add_table(
        doc,
        ["정책", "값", "이유"],
        [
            ["risk level", "LOW <20, ELEVATED 20~44, MEDIUM 45~74, HIGH 75~89, CRITICAL >=90", "대시보드와 백엔드가 같은 점수대를 해석하도록 단계화했다."],
            ["audio_only_cap", "텍스트 위험 있음 62, 텍스트 위험 없음 38", "음성만으로 100점이 찍히는 문제를 막고 영상 근거 없는 오탐을 줄인다."],
            ["weak_video_audio_cap", "텍스트 위험 있음 82, 텍스트 위험 없음 58", "영상 근거가 약한 상황에서 음성만으로 danger가 남발되지 않게 한다."],
            ["violence video-only cap", "일반 72, 강한 영상 근거 82", "폭력은 영상만으로도 감지할 수 있지만, 음성이 있을 때 더 위험하다는 차이를 보여주기 위한 상한이다."],
            ["hysteresis", "최근 5초 hit 유지, 70점 이상이면 짧은 hold", "점수가 임계값 근처에서 깜빡이며 경고가 켜졌다 꺼지는 현상을 줄인다."],
            ["cooldown", "클래스별 10~14초", "같은 사건이 너무 자주 백엔드/푸시로 반복되는 것을 막는다."],
        ],
        widths=[2.0, 2.5, 4.4],
        font_size=7.8,
    )
    add_table(
        doc,
        ["backend class", "warning", "danger", "confirm/window/cooldown", "선택 이유"],
        [
            ["violence", "58", "78", "2 hit / 8초 / 10초", "폭력은 빠른 반응이 필요하지만 단일 순간 오탐이 많아 2회 확인을 요구한다."],
            ["fall(collapse)", "52", "72", "1 hit / 8초 / 10초", "쓰러짐은 놓치면 위험하므로 confirm hit를 낮추고 danger 기준도 폭력보다 낮춘다."],
            ["loitering", "64", "82", "2 hit / 14초 / 14초", "대기/정지 오탐이 많으므로 더 긴 관찰 창과 높은 전송 기준을 둔다."],
            ["audio_risk", "68", "82", "2 hit / 20초 / 12초", "음성만으로 알림이 과도하게 발생하지 않게 높은 기준과 긴 window를 둔다."],
            ["abnormal/default", "55~60", "78~80", "2 hit / 10초 / 10초", "정체 불명 신호는 보수적으로 보내되 완전히 무시하지 않는다."],
        ],
        widths=[1.4, 0.8, 0.8, 2.3, 4.0],
        font_size=7.5,
    )

    doc.add_heading("7. WebRTC, 백엔드, 위험 클립 상호작용", level=1)
    add_image(doc, USER_VISUALS / "detectwarning_webrtc_direct_sequence_ppt_white.png", "WebRTC direct-first 시퀀스: 백엔드는 offer만 중계", max_width=7.05, max_height=4.1)
    add_image(doc, USER_VISUALS / "detectwarning_backend_event_sequence_ppt_white.png", "위험 이벤트 전송과 백엔드/PWA/푸시 흐름", max_width=7.05, max_height=4.1)
    add_table(
        doc,
        ["구성", "동작 방식", "관련 파일"],
        [
            ["CCTV 코드", "클라이언트별 영구 코드로 앱 사용자가 같은 코드를 입력하면 같은 CCTV를 본다.", "inference_server.py, camera_uploader.py"],
            ["CCTV sync", "code, name, location, status, latestRiskScore, streamUrl, inferenceStreamUrl을 /inference/cctvs로 보낸다.", "inference_server.py, app_backend_client.py"],
            ["위험 이벤트", "riskLevel, riskClass, riskScore, videoReason, audioRiskSignalDetected, clipUrl을 /inference/events로 보낸다.", "inference_server.py, backend_bridge.py"],
            ["WebRTC", "PWA/백엔드는 /api/client/{client_id}/webrtc/offer에 offer를 보내고 추론 서버는 processed frame track answer를 반환한다.", "webrtc_stream.py, inference_server.py"],
            ["클립 저장", "분석 프레임 링버퍼에서 위험 전후 MP4를 만들고 /api/events/{event_id}/clip.mp4로 제공한다.", "event_clip_service.py, inference_server.py"],
            ["Range 지원", "iPhone Safari/PWA 재생을 위해 HTTP Range 요청을 처리한다.", "inference_server.py, event_clip_service.py"],
        ],
        widths=[1.5, 5.0, 2.4],
        font_size=7.8,
    )

    doc.add_heading("8. 학습 파이프라인 상세", level=1)
    add_image(doc, USER_VISUALS / "detectwarning_inference_pipeline_ppt_white.png", "학습/추론 파이프라인 흐름도", max_width=5.5, max_height=7.2)
    add_table(
        doc,
        ["파일", "라인", "담당 로직", "상호작용/출력", "주요 클래스", "주요 함수"],
        file_detail_rows(KEY_TRAINING_FILES),
        widths=[1.35, 0.45, 2.55, 2.55, 1.25, 1.8],
        font_size=7.1,
    )
    add_table(
        doc,
        ["단계", "입력", "처리", "출력", "담당 파일"],
        [
            ["AIHub 다운로드", "datasetkey/filekey/API key", "aihubshell 다운로드, 병합, 임시파일 정리, raw manifest 생성", "current_raw/cumulative_raw", "action_training_pipeline.py"],
            ["파일키 추천", "AIHub filekey 목록, class/outside/inside 정책", "violence outside 우선, collapse/loitering outside+inside, abduction 제한/제외 판단", "추천 queue", "dashboard_aihub.py"],
            ["guideline clip", "split manifest + XML", "event start/duration, context, normal clip, filekey 추적, fallback 유지", "guideline_prepared_*.jsonl", "guideline_pose_dataset.py"],
            ["pose prepare", "clip/video", "bbox/keypoint 추출, person_not_detected fallback, pose npz 저장", "pose_path 포함 manifest", "action_training_pipeline.py"],
            ["RGB/I3D feature", "guideline manifest + source video", "clip frame sampling, feature npz, sample_weight 기록", "rgb_feature_path 포함 manifest", "extract_rgb_video_features.py"],
            ["pose 학습", "active_prepared manifest", "FocalLoss, balanced weight, MeanMax/Temporal classifier 학습", "best_action_model.pt, metrics.json", "action_model.py"],
            ["전용 task", "normal/abnormal 및 class manifest", "감지 전용과 분류 전용 task 분리", "specialized task artifacts", "specialized_action_tasks.py"],
            ["BiLSTM 미세조정", "fight/noFight 또는 fall/normal 짧은 영상", "CNN feature + BiLSTM + attention + seed sweep", "보조 booster artifacts", "train_fight_bilstm.py"],
            ["앙상블/hybrid", "pose/RGB/meta/trial checkpoint", "probability ensemble, classical hybrid search", "ensemble/hybrid summary", "ensemble_action_models.py, hybrid_pose_ensemble.py"],
            ["자동 튜닝", "기존 추출 데이터", "여러 seed/config/trial을 돌려 best metric 선택", "best config/promoted artifact", "auto_tune_action_training.py"],
        ],
        widths=[1.4, 1.6, 3.0, 2.0, 2.0],
        font_size=7.6,
    )

    doc.add_heading("9. 학습 대시보드와 운영 지원 파일", level=1)
    add_table(
        doc,
        ["파일", "라인", "담당 로직", "상호작용/출력", "주요 클래스", "주요 함수"],
        file_detail_rows(KEY_DASHBOARD_FILES),
        widths=[1.35, 0.45, 2.55, 2.55, 1.25, 1.8],
        font_size=7.1,
    )
    add_table(
        doc,
        ["기능", "관련 파일", "설명"],
        [
            ["분포 맞춰 자동 추출", "training_dashboard.py, dashboard_aihub.py", "클래스별 현재 누적 분포와 filekey 정책을 보고 다음 filekey를 추천한다."],
            ["병렬 다운로드", "training_dashboard.py, predownload_aihub_filekey.py", "작업 중 파일키와 별개로 준비 슬롯을 관리해 다음 추출 대기 시간을 줄인다."],
            ["작업 재시작", "training_dashboard.py", "pose/guideline/RGB/I3D 특정 stage부터 다시 시작하고 중복 결과는 manifest 기준으로 정리한다."],
            ["로그/진행률", "dashboard_runtime.py, reporting.py", "pipeline_status.json과 launcher log를 읽어 단계, filekey, progress를 보여준다."],
            ["품질 진단", "training_insights.py, dashboard_quality.py", "class F1, confusion matrix, validation loss, XML matching, missing source를 해석한다."],
            ["알림", "dashboard_notifications.py", "filekey stage 시작/완료/오류 시 ntfy로 요약 알림을 보낸다."],
        ],
        widths=[1.8, 2.6, 4.7],
        font_size=7.8,
    )

    doc.add_heading("10. 설정값과 산출물 기준", level=1)
    add_table(
        doc,
        ["구분", "경로", "역할"],
        [
            ["위험 음성 규칙", "configs/risk_rules.json", "A1~A10 위험 발화 카테고리, base_score, 패턴, 저위험 과장 표현 완화 규칙"],
            ["학습 config", "configs/action_training*.json", "AIHub dataset/filekey, label mapping, model/training/auto_tune 설정"],
            ["CLOVA env", "clova.env / clova.env.example", "CLOVA client id/secret 등 민감정보를 코드 밖에서 관리"],
            ["pipeline 상태", "training_data/action_pipeline_aihub/pipeline_status.json", "현재 stage/progress/message/error를 대시보드가 읽음"],
            ["manifest", "training_data/action_pipeline_aihub/manifests/*.jsonl", "raw/split/prepared/guideline/RGB feature 샘플 목록"],
            ["학습 metrics", "training_data/action_pipeline_aihub/artifacts/**/metrics.json", "accuracy, macro F1, class report, confusion matrix 저장"],
            ["모델 checkpoint", "training_data/action_pipeline_aihub/artifacts/**/best_action_model.pt", "추론 서버가 로드할 pose/action model"],
            ["위험 클립", "training_data/action_pipeline_aihub/realtime_event_clips/*.mp4", "위험 이벤트별 분석 프레임 MP4"],
        ],
        widths=[1.6, 3.2, 4.2],
        font_size=7.8,
    )

    doc.add_heading("11. 파일 간 직접 의존 관계", level=1)
    add_table(
        doc,
        ["파일", "직접 참조하는 로컬 모듈"],
        dependency_rows(KEY_RUNTIME_FILES + KEY_TRAINING_FILES + KEY_DASHBOARD_FILES),
        widths=[2.5, 6.4],
        font_size=7.1,
    )

    doc.add_heading("12. 발표/보고용 추가 시각화 자료", level=1)
    add_paragraph(
        doc,
        "아래 도표들은 발표 자료나 최종 보고서에 바로 가져다 쓰기 좋도록 크게 배치했다. "
        "같은 내용을 텍스트로 설명하는 대신, 시스템 흐름과 데이터 경계, 위험 판단 과정을 한 장씩 보여주는 용도다.",
    )
    add_image(doc, PRESENTATION_VISUALS / "02_risk_state_machine.png", "위험 판단 상태 머신: 정상, 의심, 위험 확정, cooldown 흐름", max_width=7.05, max_height=4.0)
    add_image(doc, PRESENTATION_VISUALS / "04_class_guardrails.png", "클래스별 오탐 방지 guardrail: violence/collapse/loitering별 다른 기준", max_width=7.05, max_height=4.0)
    add_image(doc, USER_VISUALS / "detectwarning_audio_flow_ppt_white.png", "음성 처리 흐름: 오디오 gate, CLOVA STT, 위험 점수화, 개인정보 경계", max_width=7.05, max_height=4.0)
    add_image(doc, USER_VISUALS / "detectwarning_risk_flow_ppt_white.png", "최종 위험 판단 흐름: 영상 점수, 음성 점수, 결합 점수, 이벤트 확정", max_width=7.05, max_height=4.0)
    add_image(doc, USER_VISUALS / "detectwarning_clip_flow_ppt_white.png", "위험 구간 클립 생성 흐름: 링버퍼, 전후 구간 추출, MP4 제공", max_width=7.05, max_height=4.0)
    add_image(doc, USER_VISUALS / "detectwarning_cctv_link_sequence_ppt_white.png", "CCTV 연결 코드 흐름: 클라이언트 고유 코드, 앱 입력, 백엔드 동기화", max_width=7.05, max_height=4.0)
    add_image(doc, USER_VISUALS / "detectwarning_sequence_ppt_white.png", "추론 서버와 백엔드/PWA의 시퀀스 흐름", max_width=7.05, max_height=4.0)
    add_image(doc, USER_VISUALS / "detectwarning_er_diagram_ppt_white.png", "앱 백엔드 연동 관점 ERD: 사용자, CCTV, 이벤트, 클립/알림 관계", max_width=7.05, max_height=4.0)

    add_landscape_section(doc)
    doc.add_heading("13. app 폴더 전체 파일 요약", level=1)
    add_paragraph(
        doc,
        "아래 표는 app 폴더의 Python 파일을 전부 훑어 AST 기준 라인 수, 주요 클래스, 주요 함수를 정리한 부록이다. "
        "상세 설명은 핵심 파일 표에 우선 배치했고, 나머지 파일은 실행 보조, 변환, 공유, 테스트, 과거 실험용 모듈로 연결된다.",
        size=8.8,
    )
    add_table(
        doc,
        ["파일", "라인", "요약", "주요 클래스", "주요 함수"],
        all_app_rows(),
        widths=[1.9, 0.45, 3.4, 2.0, 3.0],
        font_size=6.7,
    )
    add_portrait_section(doc)

    doc.add_heading("14. 테스트와 검증 기준", level=1)
    add_table(
        doc,
        ["테스트 파일", "라인", "주요 테스트 함수", "검증 의미"],
        test_rows(),
        widths=[2.6, 0.6, 3.2, 3.0],
        font_size=7.1,
    )
    add_callout(
        doc,
        "검증 원칙",
        "문서 생성은 실제 추론/학습 로직을 변경하지 않는다. 산출물 검증은 Python 문법 검사, DOCX zip 구조 검사, 포함 이미지 검사로 수행한다. 이 환경에는 LibreOffice가 없어 페이지 렌더링 QA는 별도로 수행해야 한다.",
        fill=LIGHT_BLUE,
    )

    doc.add_heading("15. 발표에서 강조할 구현 포인트", level=1)
    add_bullets(
        doc,
        [
            "음성이 없을 때의 videoOnlyScore와 음성이 있을 때의 riskScore를 함께 보여주므로, 음성 인식 도입으로 위험 감지 점수가 얼마나 상승했는지 설명할 수 있다.",
            "위험 클래스는 violence/collapse/loitering으로 직관화했고, 음성만 위험한 경우는 audio_risk 또는 의심 이벤트로 분리해 오탐을 줄였다.",
            "WebRTC direct-first 구조를 사용해 백엔드가 영상 프레임 변환 병목이 되지 않도록 했다.",
            "위험 이벤트마다 분석 프레임 기반 클립을 남겨, 왜 위험으로 판단했는지 사후 확인할 수 있다.",
            "STT 원문을 백엔드로 보내지 않기 때문에 개인정보 측면에서 더 안전한 구조다.",
            "학습 파이프라인은 AIHub guideline XML, event context, normal clip, RGB/I3D feature, pose fallback, 자동 튜닝, 앙상블까지 연결되어 있다.",
        ]
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
