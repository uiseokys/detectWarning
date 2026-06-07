from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DIAGRAMS = DOCS / "diagrams"
OUT = DOCS / "detectWarning_final_report.docx"
DOWNLOADS_ZIP = Path(r"C:\Users\Administrator\Downloads\Downloads.zip")
USER_PDF_DIR = DOCS / "_user_pdf_visuals"
USER_PDF_PNG_DIR = DOCS / "_user_pdf_visuals_png"

FONT = "Malgun Gothic"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "111827"
MUTED = "4B5563"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
LIGHT_RED = "FDE2E2"
LIGHT_YELLOW = "FFF4CE"
LIGHT_GREEN = "DCFCE7"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = INK, size: int = 9) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color, before, after in [
        ("Heading 1", 17, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("detectWarning 최종 구현 보고서")
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def configure_section(section, *, landscape: bool = False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.45)
        section.left_margin = Inches(0.45)
        section.right_margin = Inches(0.45)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)


def fit_image_inches(image_path: Path, max_width: float, max_height: float) -> tuple[float, float]:
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    aspect = max(width_px, 1) / max(height_px, 1)
    width = max_width
    height = width / aspect
    if height > max_height:
        height = max_height
        width = height * aspect
    return max(width, 0.1), max(height, 0.1)


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("detectWarning")
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run("실시간 위험상황 감지 시스템 최종 구현 보고서")
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    run.font.bold = True

    meta = [
        ("작성 목적", "프로젝트 구현 구조, 위험 판단 로직, 가중치/임계값 선택 이유, 백엔드 연동 및 푸시 알림 흐름을 발표/보고서용으로 정리"),
        ("작성일", "2026-06-02"),
        ("범위", "추론 서버, 카메라 업로더, CLOVA STT, 영상+음성 결합 위험 판단, WebRTC, 위험 클립, 앱 백엔드/PWA"),
        ("핵심 클래스", "violence / collapse / loitering, normal은 기준 상태, abduction은 현재 추론 주요 클래스에서 제외"),
    ]
    add_key_value_table(doc, meta, widths=(2.8, 13.6), fill=LIGHT_BLUE)
    add_callout(
        doc,
        "핵심 한 줄",
        "이 시스템은 영상만으로 위험도를 계산하는 데서 끝나지 않고, CLOVA STT 기반 음성 위험 신호를 함께 사용해 videoOnlyScore와 최종 riskScore의 차이를 보여준다. 따라서 발표에서는 '음성을 사용했을 때 위험 감지가 얼마나 강화되는지'를 수치로 설명할 수 있다.",
        LIGHT_GREEN,
    )


def add_key_value_table(doc: Document, rows: list[tuple[str, str]], *, widths=(3.0, 13.0), fill=LIGHT_GRAY) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_idx, (key, value) in enumerate(rows):
        set_cell_text(table.cell(row_idx, 0), key, bold=True, color=DARK_BLUE, size=9)
        set_cell_shading(table.cell(row_idx, 0), fill)
        set_cell_text(table.cell(row_idx, 1), value, size=9)
        table.cell(row_idx, 0).width = Cm(widths[0])
        table.cell(row_idx, 1).width = Cm(widths[1])
    doc.add_paragraph()


def add_table(doc: Document, headers: list[str], rows: list[list[str]], *, widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.cell(0, i), header, bold=True, color=DARK_BLUE, size=8.7)
        set_cell_shading(table.cell(0, i), LIGHT_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=8.2)
            if widths and i < len(widths):
                cells[i].width = Cm(widths[i])
    doc.add_paragraph()


def add_callout(doc: Document, title: str, body: str, fill: str = LIGHT_YELLOW) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(body)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(INK)
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        run.font.name = FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        run.font.size = Pt(10)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        run.font.name = FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        run.font.size = Pt(10)


def add_figure(doc: Document, image_name: str, caption: str, *, width_in: float = 6.45) -> None:
    path = DIAGRAMS / image_name
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def add_image_path(doc: Document, image_path: Path, caption: str, *, width_in: float = 6.45) -> None:
    if not image_path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def add_fitted_image_path(
    doc: Document,
    image_path: Path,
    caption: str,
    *,
    max_width_in: float,
    max_height_in: float,
) -> None:
    if not image_path.exists():
        return
    width_in, height_in = fit_image_inches(image_path, max_width_in, max_height_in)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width_in), height=Inches(height_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(4)
    r = cap.add_run(caption)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def convert_user_pdf_visuals() -> list[tuple[Path, str]]:
    """Extract PDFs from the user-provided ZIP and render their first pages as PNG."""
    existing = sorted(USER_PDF_PNG_DIR.glob("*.png")) if USER_PDF_PNG_DIR.exists() else []
    if existing:
        return [(path, f"{path.stem}.pdf") for path in existing]
    if not DOWNLOADS_ZIP.exists():
        return []
    try:
        import pypdfium2 as pdfium
    except Exception:
        return []

    USER_PDF_DIR.mkdir(parents=True, exist_ok=True)
    USER_PDF_PNG_DIR.mkdir(parents=True, exist_ok=True)
    converted: list[tuple[Path, str]] = []
    with zipfile.ZipFile(DOWNLOADS_ZIP, "r") as archive:
        for info in sorted(archive.infolist(), key=lambda item: item.filename):
            if not info.filename.lower().endswith(".pdf"):
                continue
            pdf_path = USER_PDF_DIR / Path(info.filename).name
            with archive.open(info) as source, pdf_path.open("wb") as target:
                target.write(source.read())
            png_path = USER_PDF_PNG_DIR / f"{pdf_path.stem}.png"
            pdf = pdfium.PdfDocument(str(pdf_path))
            try:
                if len(pdf) <= 0:
                    continue
                page = pdf[0]
                bitmap = page.render(scale=2.0)
                pil_image = bitmap.to_pil()
                if pil_image.mode != "RGB":
                    pil_image = pil_image.convert("RGB")
                pil_image.save(png_path, quality=95)
            finally:
                pdf.close()
            converted.append((png_path, pdf_path.name))
    return converted


def visual_caption_from_filename(filename: str) -> str:
    mapping = {
        "detectwarning_architecture_ppt_white.pdf": "사용자 제공 PDF 도표: 전체 시스템 아키텍처",
        "detectwarning_audio_flow_ppt_white.pdf": "사용자 제공 PDF 도표: 음성 인식 및 위험 신호 흐름",
        "detectwarning_backend_event_sequence_ppt_white.pdf": "사용자 제공 PDF 도표: 백엔드 이벤트 저장 및 푸시 시퀀스",
        "detectwarning_cctv_link_sequence_ppt_white.pdf": "사용자 제공 PDF 도표: CCTV 코드 연결 시퀀스",
        "detectwarning_clip_flow_ppt_white.pdf": "사용자 제공 PDF 도표: 위험 구간 클립 생성 흐름",
        "detectwarning_er_diagram_ppt_white.pdf": "사용자 제공 PDF 도표: 앱 백엔드 ERD",
        "detectwarning_inference_pipeline_ppt_white.pdf": "사용자 제공 PDF 도표: 추론 파이프라인",
        "detectwarning_risk_flow_ppt_white.pdf": "사용자 제공 PDF 도표: 위험 점수 판단 흐름",
        "detectwarning_sequence_ppt_white.pdf": "사용자 제공 PDF 도표: 전체 실시간 시퀀스",
        "detectwarning_webrtc_direct_sequence_ppt_white.pdf": "사용자 제공 PDF 도표: WebRTC direct-first 시퀀스",
    }
    return mapping.get(filename, f"사용자 제공 PDF 도표: {filename}")


def add_user_pdf_figure(
    doc: Document,
    visual_map: dict[str, Path],
    source_name: str,
    caption: str,
    *,
    width_in: float = 6.45,
) -> None:
    image_path = visual_map.get(source_name)
    if image_path and image_path.exists():
        add_image_path(doc, image_path, f"{caption} / 원본: Downloads.zip::{source_name}", width_in=width_in)
        return
    add_callout(
        doc,
        "원본 PDF 도표 누락",
        f"Downloads.zip에서 {source_name} 변환 이미지를 찾지 못했다. 문서 생성 전에 PDF 변환 단계가 정상 완료됐는지 확인해야 한다.",
        LIGHT_RED,
    )


def add_user_pdf_visual_page(
    doc: Document,
    visual_map: dict[str, Path],
    title: str,
    source_name: str,
    caption: str,
    *,
    landscape: bool,
) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section, landscape=landscape)
    heading = doc.add_heading(title, level=1)
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(6)
    image_path = visual_map.get(source_name)
    if image_path and image_path.exists():
        add_fitted_image_path(
            doc,
            image_path,
            f"{caption} / 원본: Downloads.zip::{source_name}",
            max_width_in=9.9 if landscape else 7.1,
            max_height_in=6.55 if landscape else 8.65,
        )
    else:
        add_callout(
            doc,
            "원본 PDF 도표 누락",
            f"Downloads.zip에서 {source_name} 변환 이미지를 찾지 못했다. 문서 생성 전에 PDF 변환 단계가 정상 완료됐는지 확인해야 한다.",
            LIGHT_RED,
        )
    next_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(next_section, landscape=False)


def add_section_break(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section, landscape=False)


def build_report() -> None:
    user_pdf_visuals = convert_user_pdf_visuals()
    user_visual_map = {source_name: image_path for image_path, source_name in user_pdf_visuals}
    doc = Document()
    style_document(doc)
    add_title(doc)

    doc.add_heading("1. 프로젝트 목적과 설계 방향", level=1)
    doc.add_paragraph(
        "detectWarning은 CCTV, 웹캠, 노트북 카메라, 영상 테스트 파일을 입력으로 받아 위험행동과 위험음성을 동시에 분석하는 실시간 위험상황 감지 시스템이다. "
        "초기에는 pose 기반 행동 분류만으로 시작했지만, 단일 영상 모델만으로는 앉아 있음과 쓰러짐, 기다림과 배회, 장난스러운 움직임과 폭력의 경계를 안정적으로 나누기 어려웠다. "
        "따라서 최종 구조는 영상 기반 행동분석, CLOVA STT 기반 음성 위험 신호, 시간 안정화, 백엔드 이벤트 정책을 결합하는 방식으로 설계했다."
    )
    add_bullets(
        doc,
        [
            "실사용 목표는 단순 분류 정확도보다 위험을 놓치지 않으면서 오탐을 줄이는 것이다.",
            "발표 관점의 핵심 강점은 영상만 점수와 영상+음성 점수를 동시에 보여줘 음성 인식의 기여도를 설명할 수 있다는 점이다.",
            "현재 위험 클래스는 violence, collapse, loitering 3개이며, normal은 정상 기준 상태로 사용한다.",
            "abduction은 현재 최종 추론 로직에서 주요 위험 클래스로 쓰지 않도록 정리했다.",
        ],
    )
    add_callout(
        doc,
        "왜 단순 OR 결합을 쓰지 않았는가",
        "영상 또는 음성 중 하나라도 위험이면 바로 danger로 올리는 방식은 빠르지만 오탐이 많다. 그래서 이 프로젝트는 videoOnlyScore, audioScore, fusion bonus, 클래스별 threshold, hysteresis를 조합한다. 음성만으로는 의심 단계에 머물게 하고, 영상과 음성이 같은 사건을 지지할 때 점수를 더 크게 올리도록 했다.",
    )

    add_user_pdf_visual_page(
        doc,
        user_visual_map,
        "2. 전체 시스템 아키텍처",
        "detectwarning_architecture_ppt_white.pdf",
        "그림 1. detectWarning 전체 시스템 아키텍처",
        landscape=True,
    )
    doc.add_paragraph(
        "시스템은 크게 카메라/영상 입력, 추론 서버, 앱 백엔드, PWA 화면으로 나뉜다. 추론 서버는 무거운 AI 분석을 담당하고, 앱 백엔드는 사용자 인증, CCTV 코드 연결, 이벤트 저장, 푸시 알림, PWA 화면을 담당한다. "
        "이렇게 분리한 이유는 AI 추론과 사용자 서비스의 장애 범위를 분리하고, STT 원문 같은 민감 정보를 앱 백엔드로 넘기지 않기 위해서다."
    )
    add_table(
        doc,
        ["구성", "담당 역할", "분리한 이유"],
        [
            ["카메라 업로더", "Mac/Windows/iPhone 카메라 프레임과 음성 조각 전송", "여러 클라이언트를 CCTV처럼 붙일 수 있게 한다."],
            ["추론 서버", "YOLO/Pose/RGB-I3D/BiLSTM/STT/RiskAnalyzer 수행", "GPU를 쓰는 무거운 처리를 한 서버에 집중한다."],
            ["앱 백엔드", "회원, CCTV 코드, 위험 이벤트, 푸시, PWA 제공", "사용자 서비스와 데이터 저장을 안정적으로 관리한다."],
            ["PWA/iPhone", "연결 CCTV 조회, 실시간 영상, 위험 기록, 푸시 수신", "사용자에게 필요한 정보만 보여준다."],
        ],
        widths=[3.0, 6.5, 6.0],
    )

    doc.add_heading("3. 주요 구현 파일과 책임", level=1)
    add_table(
        doc,
        ["파일", "책임", "설계 포인트"],
        [
            ["app/inference_server.py", "FastAPI 서버, 대시보드, 클라이언트 세션, 분석 API, 백엔드 전송", "큰 파일이지만 현재는 동작 안정성을 위해 라우팅 중심으로 유지한다."],
            ["app/action_realtime.py", "실시간 행동 인식, RGB/I3D, pose 보조, fight/fall BiLSTM 점수", "단일 프레임이 아니라 최근 프레임 버퍼를 사용한다."],
            ["app/risk_analyzer.py", "영상 점수, 음성 점수, 결합 점수, 위험 단계 산정", "가중치와 오탐 억제 정책의 중심 파일이다."],
            ["app/stt_service.py", "CLOVA CSR 호출, 음성 보정, gate, rate limit", "음성 원문은 추론 서버 내부에서만 사용한다."],
            ["app/webrtc_stream.py", "분석 프레임을 WebRTC video track으로 송출", "JPEG polling 대신 direct-first 구조로 FPS 저하를 줄인다."],
            ["app/event_clip_service.py", "위험 구간 분석 프레임 링버퍼와 MP4 클립 생성", "전 5초+후 10초 클립으로 사건 근거를 남긴다."],
            ["app/backend_bridge.py", "백엔드 전송 circuit breaker", "백엔드 장애가 추론 서버를 멈추지 않게 한다."],
        ],
        widths=[4.0, 6.0, 6.0],
    )

    add_user_pdf_visual_page(
        doc,
        user_visual_map,
        "4. 실시간 감지 처리 흐름",
        "detectwarning_sequence_ppt_white.pdf",
        "그림 2. 실시간 감지 시퀀스",
        landscape=True,
    )
    add_numbered(
        doc,
        [
            "카메라 업로더 또는 영상 테스트 클라이언트가 프레임을 추론 서버로 보낸다.",
            "추론 서버는 사람 탐지, pose overlay, 행동 분류, 최신 분석 프레임 저장을 수행한다.",
            "음성 조각이 있으면 STT 큐에 넣고, CLOVA CSR 결과를 위험음성 분석에 전달한다.",
            "RiskAnalyzer가 영상 점수, 음성 점수, 결합 점수, 최근 이력, 클래스별 정책을 합산한다.",
            "백엔드 전송 정책을 통과한 사건만 /inference/events로 전달한다.",
            "백엔드는 사건을 저장하고 SSE 및 Web Push로 PWA 사용자에게 알린다.",
            "위험 이벤트가 발생하면 분석 프레임 링버퍼에서 MP4 클립을 만든다.",
        ],
    )

    add_user_pdf_visual_page(
        doc,
        user_visual_map,
        "5. 영상 행동 분석 로직",
        "detectwarning_inference_pipeline_ppt_white.pdf",
        "그림 3. 추론 파이프라인",
        landscape=False,
    )
    doc.add_paragraph(
        "영상 분석은 YOLO 사람 탐지, pose/관절점 기반 보조 정보, RGB/I3D 계열 특징, BiLSTM 보조 점수를 함께 사용한다. "
        "특히 violence와 collapse는 짧은 순간의 움직임이 중요하고, loitering은 시간 지속성이 중요하기 때문에 클래스별로 서로 다른 안정화 기준을 적용한다."
    )
    add_table(
        doc,
        ["변수", "값", "사용 이유"],
        [
            ["clip_seconds", "4.0초", "너무 짧으면 행동 변화가 보이지 않고, 너무 길면 반응성이 떨어진다. 4초는 실시간성과 행동 맥락의 균형점이다."],
            ["min_interval_seconds", "1.5초", "매 프레임마다 무거운 행동 모델을 돌리지 않고 GPU/CPU 여유를 남긴다."],
            ["sequence_length", "32", "pose 시퀀스가 너무 짧아지는 것을 막고, 쓰러짐/배회 같은 시간 패턴을 보게 한다."],
            ["rgb_frames", "16", "2D/3D feature 추출에 필요한 대표 프레임 수를 확보하면서 지연을 제한한다."],
            ["image_size", "112", "실시간 추론에서 속도를 우선하면서 행동 특징은 유지하는 입력 크기다."],
            ["normal_threshold", "0.78", "정상 상태를 강하게 유지해 약한 이상 점수로 바로 위험이 뜨는 것을 줄인다."],
            ["min_action_confidence", "0.45", "모델이 완전히 불확실한 결과를 버리되, 보조 신호와 결합될 여지는 남긴다."],
        ],
        widths=[3.5, 2.3, 10.0],
    )
    add_table(
        doc,
        ["클래스", "주요 조건", "선택 이유"],
        [
            ["violence", "min_confidence 0.74, min_abnormal 0.86, min_streak 2, fast_accept 0.97/0.90", "폭력은 반응성이 중요하지만 오탐도 많다. 그래서 강한 영상 신호는 빠르게 받아들이고, 일반 신호는 2회 이상 반복을 요구한다."],
            ["collapse", "min_confidence 0.58, min_abnormal 0.76, static_abnormal 0.90, temporal 0.94", "쓰러짐은 confidence가 낮게 나올 수 있어 문턱을 낮추되, 단일 프레임 오탐은 시간 안정화와 자세 근거로 억제한다."],
            ["loitering", "min_confidence 0.68, min_abnormal 0.78, min_streak 3, min_duration 8초", "기다리거나 서 있는 사람을 배회로 잘못 잡지 않도록 반복성과 지속시간을 강하게 요구한다."],
        ],
        widths=[2.4, 6.3, 8.0],
    )
    add_callout(
        doc,
        "loitering 오탐을 줄인 방식",
        "배회는 단순히 같은 위치에 오래 서 있는 것과 구분해야 한다. 그래서 평균 이동량 6.0, 중심 이동 45.0, 정지 프레임 12개 기준을 함께 보고, 기다림 패턴이면 점수를 0.45배로 줄이거나 34점 이하로 제한한다.",
    )

    add_user_pdf_visual_page(
        doc,
        user_visual_map,
        "6. 위험 점수 산정 로직",
        "detectwarning_risk_flow_ppt_white.pdf",
        "그림 4. 위험 점수 산정 로직",
        landscape=False,
    )
    doc.add_paragraph(
        "RiskAnalyzer는 최종 점수를 하나의 모델 confidence가 아니라 여러 점수의 합으로 만든다. 기본 구조는 영상 컴포넌트, 음성 컴포넌트, 영상+음성 결합 보너스, hysteresis를 거쳐 최종 riskScore를 만드는 방식이다."
    )
    add_table(
        doc,
        ["점수", "의미", "발표에서 설명할 포인트"],
        [
            ["videoOnlyScore", "영상만으로 계산한 위험도", "음성을 꺼도 어느 정도 위험을 잡는지 보여준다."],
            ["audioScore", "STT/소리 패턴만으로 계산한 위험 신호", "음성이 단독으로 얼마나 위험 신호를 주는지 확인한다."],
            ["audioVideoGain", "음성 때문에 올라간 점수", "교수님이 말한 음성 사용 효과를 숫자로 보여주는 핵심 값이다."],
            ["riskScore", "최종 결합 점수", "백엔드 이벤트, 푸시 알림, 위험 기록의 기준이 된다."],
        ],
        widths=[3.2, 6.0, 7.0],
    )
    add_table(
        doc,
        ["단계", "점수 기준", "의미"],
        [
            ["LOW", "0~19", "정상 또는 의미 없는 신호"],
            ["ELEVATED", "20~44", "약한 이상 신호, 대시보드 관찰용"],
            ["MEDIUM", "45~74", "주의 후보, 반복/정책 확인 후 백엔드 전송 가능"],
            ["HIGH", "75~89", "위험 후보, 백엔드 danger 매핑 가능"],
            ["CRITICAL", "90~100", "강한 위험, 기존 사건도 재알림할 수 있는 수준"],
        ],
        widths=[2.5, 3.0, 10.5],
    )
    add_callout(
        doc,
        "45 / 75 / 90 기준을 둔 이유",
        "45점은 모델이 약하게 의심하는 구간과 실제 주의 상황을 나누는 최소 경계다. 75점은 영상 또는 음성+영상 근거가 충분히 강한 위험 경계다. 90점은 사람에게 즉시 알려야 하는 매우 강한 신호로, critical 단계와 재알림 조건에 사용한다.",
    )

    doc.add_heading("7. 클래스별 가중치와 선택 이유", level=1)
    add_table(
        doc,
        ["클래스", "주요 가중치/임계값", "왜 이렇게 설정했는가"],
        [
            ["violence", "base 18, confidence weight 20, abnormal weight 14, audio bonus 16, multi-person +4, video-only cap 72/82", "폭력은 영상 confidence가 높아도 장난/빠른 움직임 오탐이 있다. 그래서 기본 점수는 낮게 두고 confidence와 abnormal을 반영하되, 음성과 다중 사람 근거가 있을 때 강하게 올린다."],
            ["collapse", "audio bonus 14, min confidence 0.58, min abnormal 0.76, temporal abnormal 0.94, stable +5", "쓰러짐은 실제 상황에서 사람 일부만 보이거나 confidence가 낮을 수 있다. 누락을 줄이기 위해 confidence 문턱은 낮추고, 대신 시간/자세 확인으로 단일 프레임 오탐을 줄인다."],
            ["loitering", "audio bonus 8, min confidence 0.68, min abnormal 0.78, confirm window 12초, min duration 8초, video-only cap 62", "배회는 기다림과 오탐이 가장 잘 섞인다. 그래서 점수 상승을 보수적으로 하고, 영상만으로는 62점 근처에서 제한해 바로 danger로 튀지 않게 했다."],
        ],
        widths=[2.2, 6.7, 8.0],
    )
    add_table(
        doc,
        ["보정 정책", "값", "효과"],
        [
            ["weak_signal_multiplier", "violence 0.42 / collapse 0.58 / loitering 0.50", "confidence 또는 abnormal_score가 낮은 애매한 신호의 점수를 줄인다."],
            ["stable_bonus", "violence 6 / collapse 5 / loitering 6", "최근 같은 클래스가 반복되면 단일 순간이 아니라 사건으로 판단한다."],
            ["audio_only_cap", "텍스트 위험 있음 62 / 텍스트 없음 38", "음성만으로 바로 100점이 되던 문제를 막는다."],
            ["weak_video_audio_cap", "텍스트 위험 있음 82 / 텍스트 없음 58", "영상 근거가 약한데 소리만 큰 경우를 과하게 위험으로 보내지 않는다."],
            ["risk hysteresis", "5초 내 45점 이상 반복 또는 70점 이상이면 3초 hold", "점수가 깜빡이는 현상을 줄이고, 위험 표시가 너무 빨리 사라지지 않게 한다."],
        ],
        widths=[4.5, 4.5, 7.5],
    )

    add_user_pdf_visual_page(
        doc,
        user_visual_map,
        "8. 음성 인식과 개인정보 보호",
        "detectwarning_audio_flow_ppt_white.pdf",
        "그림 5. CLOVA 음성 인식 및 개인정보 흐름",
        landscape=True,
    )
    doc.add_paragraph(
        "음성은 프로젝트의 차별점이지만 개인정보 위험도 크다. 그래서 STT 원문, matched keywords, 대화 내용은 앱 백엔드로 보내지 않고 추론 서버 내부에서만 위험 신호 계산에 사용한다. 백엔드에는 audioRiskSignalDetected, audioScore, audioVideoGain, audioEventClasses 같은 메타데이터만 전송한다."
    )
    add_table(
        doc,
        ["항목", "값/정책", "이유"],
        [
            ["음성 후보", "raw / normalized / boosted", "촬영 환경마다 음량이 다르므로 CLOVA에 여러 보정 후보를 시도해 인식률을 높인다."],
            ["normalized RMS", "0.075, max gain 5.0", "작은 목소리를 보정하되 과증폭으로 잡음이 커지는 것을 제한한다."],
            ["boosted RMS", "0.11, max gain 4.0", "특정 테스트 영상처럼 소리가 낮게 들어오는 경우의 백업 후보다."],
            ["최소 길이", "0.55초", "너무 짧은 조각은 CLOVA 비용만 늘리고 의미 있는 문장이 되기 어렵다."],
            ["최소 RMS", "0.006", "거의 무음인 구간을 STT로 보내지 않는다."],
            ["최소 peak", "0.035", "RMS는 낮아도 순간 발화가 있는지 확인한다."],
            ["active ratio", "0.018", "전체 구간에서 실제 음성 활동이 너무 적으면 건너뛴다."],
            ["업로드 제한", "3초 간격, 분당 18초", "CLOVA 비용 폭증과 대기열 적체를 막는다."],
        ],
        widths=[3.5, 4.5, 8.3],
    )
    add_callout(
        doc,
        "왜 감지된 소리를 그대로 보내지 않았는가",
        "휴대폰 영상이나 노트북 마이크는 음량, 잡음, 좌우 채널, 피크가 제각각이다. 그대로 보내면 어떤 영상은 너무 작게 들어가고 어떤 영상은 클리핑된다. 그래서 DC offset 제거, RMS 보정, peak clipping 방지, mono 변환을 적용한 후보를 만들었다.",
    )

    doc.add_heading("9. 영상+음성 결합 판단", level=1)
    doc.add_paragraph(
        "영상과 음성은 단순히 더하는 것이 아니라 상황별로 다르게 결합한다. violence는 비명, 도움 요청, 충격음과 강하게 결합하고, collapse는 통증/도움 요청/충격음과 결합한다. loitering은 배회 자체가 음성 위험과 직접 연결되지 않는 경우가 많기 때문에 음성 bonus를 작게 두고, distress_speech처럼 명확한 위험 발화가 있을 때만 보강한다."
    )
    add_table(
        doc,
        ["상황", "결합 방식", "효과"],
        [
            ["영상 strong + 음성 strong", "class audio bonus + fusion bonus + audio+video_confirmed", "같은 사건을 두 센서가 지지하므로 최종 점수를 크게 올린다."],
            ["영상 weak + 음성 strong", "weak_video_audio_cap 적용", "위험 의심은 보이지만 영상 근거가 약하므로 위험 단계 폭주를 막는다."],
            ["영상 없음 + 음성 strong", "audio_only_cap 적용", "구조 요청 등은 기록하지만 바로 최고 위험으로 만들지 않는다."],
            ["영상 strong + 음성 없음", "video_only_cap 적용", "폭력 오탐이 너무 커지는 것을 막고, 음성이 있을 때 차이가 보이게 한다."],
        ],
        widths=[4.2, 6.0, 6.0],
    )

    add_user_pdf_visual_page(
        doc,
        user_visual_map,
        "10. CCTV 코드와 앱 백엔드 연결",
        "detectwarning_cctv_link_sequence_ppt_white.pdf",
        "그림 6. 카메라별 CCTV 코드 연결",
        landscape=True,
    )
    doc.add_paragraph(
        "각 카메라/클라이언트는 고유 CCTV 코드를 가진다. 사용자는 PWA에서 이 코드를 입력해 CCTV를 연결하고, 같은 코드를 입력한 사용자는 동등한 권한으로 해당 CCTV를 볼 수 있다. 기존 서버 전체 1개 페어링 코드 방식은 다중 카메라에 맞지 않아, 클라이언트별 고유 코드 방식으로 정리했다."
    )
    add_bullets(
        doc,
        [
            "같은 Mac/Windows 클라이언트는 안정적인 client_id를 쓰면 코드가 유지된다.",
            "추론 서버는 /inference/cctvs로 CCTV 상태, 이름, 위치, 최신 점수를 백엔드에 동기화한다.",
            "STT 원문은 동기화 payload에 포함하지 않는다.",
            "백엔드는 code 기준 upsert로 같은 CCTV를 새로 만들지 않고 갱신한다.",
        ],
    )

    add_user_pdf_visual_page(
        doc,
        user_visual_map,
        "11. WebRTC direct-first 영상 송출",
        "detectwarning_webrtc_direct_sequence_ppt_white.pdf",
        "그림 7. WebRTC direct-first 송출 구조",
        landscape=True,
    )
    add_callout(
        doc,
        "왜 JPEG polling을 피했는가",
        "JPEG를 매번 HTTP로 요청하면 백엔드가 프레임 중계 서버처럼 동작해 CPU, 네트워크, 지연이 모두 증가한다. WebRTC direct-first는 백엔드가 offer/answer만 중계하고 실제 분석 프레임은 추론 서버가 직접 video track으로 보내기 때문에 FPS와 반응성이 좋아진다.",
    )
    add_table(
        doc,
        ["설계", "값/정책", "이유"],
        [
            ["WebRTC track fps", "analysis_fps 기준, 최대 60fps clamp", "대시보드 분석 프레임과 PWA FPS 차이를 줄인다."],
            ["stale frame hold", "최근 정상 프레임을 최대 5초 유지", "짧은 네트워크 흔들림으로 화면이 바로 검게 변하지 않게 한다."],
            ["codec preference", "H264 우선, VP8 fallback", "iPhone/PWA 호환성과 브라우저 호환성을 함께 고려한다."],
            ["frame resize", "max edge 제한", "너무 큰 원본 영상이 WebRTC 인코딩 병목을 만들지 않게 한다."],
        ],
        widths=[3.5, 5.0, 8.0],
    )

    add_user_pdf_visual_page(
        doc,
        user_visual_map,
        "12. 위험 구간 클립 저장",
        "detectwarning_clip_flow_ppt_white.pdf",
        "그림 8. 위험 구간 클립 저장 파이프라인",
        landscape=False,
    )
    doc.add_paragraph(
        "위험 이벤트가 발생하면 관절점/박스가 그려진 분석 프레임 기준으로 MP4 클립을 만든다. 원본 대화나 STT 원문은 저장하지 않고, 발표나 앱 기록에서는 위험 판단 근거가 되는 분석 화면만 재생한다."
    )
    add_table(
        doc,
        ["항목", "값", "이유"],
        [
            ["링버퍼", "최근 30초, 최대 1200프레임", "판단 전 구간을 확보하면서 메모리 사용량을 제한한다."],
            ["클립 범위", "판단 전 5초 + 판단 후 10초", "위험 발생 전 맥락과 이후 확인 장면을 함께 보여준다."],
            ["클립 FPS", "최대 20fps, 최소 재생 15fps", "재생은 부드럽게 유지하되 파일 크기와 인코딩 시간을 제한한다."],
            ["Range 요청", "지원", "iPhone Safari/PWA에서 MP4 seek와 스트리밍 재생이 가능하게 한다."],
        ],
        widths=[3.2, 4.2, 8.5],
    )

    add_user_pdf_visual_page(
        doc,
        user_visual_map,
        "13. 앱 백엔드 ERD와 데이터 흐름",
        "detectwarning_er_diagram_ppt_white.pdf",
        "그림 9. 앱 백엔드 ERD",
        landscape=False,
    )
    doc.add_paragraph(
        "백엔드는 사용자, CCTV, 사용자-CCTV 연결, 위험 이벤트, 점수 샘플, 상태 이벤트, 푸시 구독을 저장한다. 핵심은 CCTV code를 기준으로 여러 사용자가 같은 CCTV에 연결될 수 있고, 위험 이벤트는 CCTV 단위로 저장된다는 점이다."
    )
    add_table(
        doc,
        ["테이블", "역할", "설계 이유"],
        [
            ["USERS / SESSIONS / USER_PROFILES", "회원가입, 로그인, 사용자 정보", "PWA 사용자 권한과 푸시 대상을 관리한다."],
            ["CCTVS", "CCTV 코드, 이름, 위치, 상태, 최신 점수", "앱 사용자가 입력하는 code와 추론 서버 client_id를 연결한다."],
            ["USER_CCTVS", "사용자와 CCTV의 다대다 연결", "같은 코드를 입력한 모든 사용자가 동등한 권한을 갖게 한다."],
            ["RISK_EVENTS", "위험 사건, 점수, 이유, clipUrl, incident_key", "기록 목록과 상세 화면, 푸시 알림의 근거가 된다."],
            ["CCTV_SCORE_SAMPLES", "시간별 위험 점수 추이", "PWA의 위험 점수 그래프를 만든다."],
            ["PUSH_SUBSCRIPTIONS", "Web Push endpoint와 인증 정보", "연결된 사용자에게 위험 알림을 보낸다."],
        ],
        widths=[4.0, 5.5, 7.0],
    )

    add_user_pdf_visual_page(
        doc,
        user_visual_map,
        "14. 백엔드 이벤트 저장 및 푸시 알림",
        "detectwarning_backend_event_sequence_ppt_white.pdf",
        "그림 10. 백엔드 이벤트 저장 및 푸시 알림 흐름",
        landscape=True,
    )
    add_table(
        doc,
        ["클래스", "warning_min", "danger_min", "confirm_hits/window", "cooldown", "이유"],
        [
            ["violence", "58", "78", "2회 / 8초", "10초", "폭력은 빠르게 알려야 하지만 단일 오탐이 많아 2회 확인을 둔다."],
            ["fall", "52", "72", "1회 / 8초", "10초", "쓰러짐은 놓치면 위험하므로 warning 문턱과 반복 요구를 낮춘다."],
            ["loitering", "64", "82", "2회 / 14초", "14초", "배회는 오탐이 많아 문턱과 확인 시간을 높인다."],
            ["audio_risk", "68", "82", "2회 / 20초", "12초", "음성 단독은 개인정보/오탐 가능성이 있어 더 보수적으로 전송한다."],
            ["abnormal", "60", "80", "2회 / 10초", "10초", "지원되지 않은 이상 신호는 중간 기준으로 처리한다."],
        ],
        widths=[2.4, 2.4, 2.4, 3.5, 2.2, 6.0],
    )
    add_callout(
        doc,
        "백엔드 전송 최적화",
        "추론 서버는 APP_BACKEND_POST_QUEUE와 circuit breaker를 사용한다. 백엔드 전송이 3회 실패하면 5초 동안 회로를 열어 불필요한 재시도를 줄인다. 이는 백엔드가 잠시 느려지거나 꺼져도 추론 서버의 영상 분석이 멈추지 않게 하기 위한 안정화 장치다.",
        LIGHT_BLUE,
    )

    doc.add_heading("15. API 요약", level=1)
    add_table(
        doc,
        ["서버", "Method/Path", "역할"],
        [
            ["추론", "GET /health", "GPU/RAM/STT/WebRTC/백엔드 연결 상태 확인"],
            ["추론", "GET /api/clients", "연결된 카메라/영상 테스트 클라이언트 목록"],
            ["추론", "POST /api/client/{client_id}/webrtc/offer", "분석 프레임 WebRTC answer 생성"],
            ["추론", "POST /analyze/frame", "카메라 프레임 분석"],
            ["추론", "POST /analyze/audio", "음성 조각 분석"],
            ["추론", "GET /api/events/{event_id}/clip.mp4", "위험 구간 MP4 클립 제공"],
            ["백엔드", "POST /inference/cctvs", "CCTV code/name/location/status 동기화"],
            ["백엔드", "POST /inference/events", "위험 이벤트 수신 및 저장"],
            ["백엔드", "GET /cctvs/{id}/events", "PWA 위험 기록 목록 조회"],
            ["백엔드", "POST /push/subscribe", "PWA Web Push 구독 저장"],
        ],
        widths=[2.2, 5.8, 8.0],
    )

    doc.add_heading("16. 발표 모드와 음성 효과 설명 방식", level=1)
    doc.add_paragraph(
        "교수님이 강조한 부분은 음성 인식을 사용했을 때 위험 인식이 얼마나 좋아지는지다. 그래서 대시보드 발표 모드에는 영상만 점수, 음성 점수, 영상+음성 결합 점수, Audio gain을 함께 표시하도록 구성했다."
    )
    add_table(
        doc,
        ["발표 지표", "계산 방식", "말할 수 있는 설명"],
        [
            ["영상만", "videoOnlyScore", "카메라 화면만 봤을 때 시스템이 판단한 위험도입니다."],
            ["음성", "audioScore", "도움 요청, 위협 표현, 큰 소리, 충격음 같은 음성 위험 신호입니다."],
            ["영상+음성", "riskScore", "영상과 음성을 함께 보았을 때의 최종 위험도입니다."],
            ["Audio gain", "rawScore - videoOnlyScore", "음성 신호가 들어오면서 위험 판단이 얼마나 강화됐는지 보여줍니다."],
        ],
        widths=[3.0, 4.2, 9.0],
    )
    add_callout(
        doc,
        "발표 예시 문장",
        "예를 들어 영상만으로는 62점이라 warning 후보였지만, 피해자의 '하지 마세요', '아파요', '경찰 불러주세요' 같은 발화가 들어오면 음성 위험 신호가 결합되어 최종 점수가 80점대로 올라갑니다. 이처럼 음성은 영상의 애매한 판단을 보강해 위험 상황을 더 빠르고 강하게 잡도록 설계했습니다.",
        LIGHT_GREEN,
    )

    doc.add_heading("17. 보안과 개인정보 정책", level=1)
    add_table(
        doc,
        ["정책", "적용 방식", "이유"],
        [
            ["STT 원문 비전송", "백엔드 payload에는 transcript/matchedKeywords 제외", "사용자의 실제 대화 내용이 앱 DB에 저장되지 않게 한다."],
            ["메타데이터만 전송", "audioRiskSignalDetected, audioScore, audioEventClasses", "위험 판단 근거는 제공하되 개인정보 노출을 줄인다."],
            ["추론 서버 token", "X-Inference-Token", "앱 백엔드 inference API를 외부 임의 호출에서 보호한다."],
            ["CCTV code", "6자리 영문+숫자 코드", "PWA 사용자가 쉽게 입력하되 내부 client_id와 분리한다."],
            ["클립 저장 범위 제한", "위험 전후 15초 분석 프레임", "필요 이상의 장시간 녹화를 피하고 사건 근거만 남긴다."],
        ],
        widths=[3.5, 5.8, 7.0],
    )

    doc.add_heading("18. 운영 안정화와 성능 최적화", level=1)
    add_bullets(
        doc,
        [
            "영상은 WebRTC direct-first로 보내 백엔드가 JPEG 중계 병목이 되지 않게 했다.",
            "STT는 별도 큐와 rate limit을 사용해 CLOVA 요청 비용과 지연을 제한했다.",
            "백엔드 전송은 circuit breaker를 사용해 장애 시 추론 루프가 멈추지 않게 했다.",
            "분석 프레임은 최신 프레임 캐시와 stale frame hold를 사용해 짧은 끊김에도 화면이 바로 사라지지 않게 했다.",
            "위험 이벤트는 incident key와 cooldown으로 같은 사건이 무한히 중복 저장되지 않게 했다.",
            "GPU 여유가 있을 때는 분석 FPS와 모델 입력 품질을 올릴 수 있지만, 발표 안정성을 위해 프레임 드롭과 STT 큐 상태를 같이 봐야 한다.",
        ],
    )

    doc.add_heading("19. 테스트와 검증 방법", level=1)
    add_table(
        doc,
        ["검증 항목", "확인 방법", "통과 기준"],
        [
            ["서버 상태", "GET /health 또는 대시보드 상태 카드", "GPU/RAM/STT/WebRTC/backend 상태가 정상 또는 명확한 오류로 표시"],
            ["카메라 연결", "GET /api/clients", "client_id, CCTV code, has_frame, last_seen_seconds 확인"],
            ["WebRTC", "대시보드/PWA 실시간 분석 프레임", "JPEG polling 없이 영상이 안정적으로 표시"],
            ["음성", "대시보드 audio status와 audioLevel", "clova_empty 또는 missing이 지속되지 않고 위험 발화가 score에 반영"],
            ["위험 이벤트", "백엔드 기록 목록과 푸시", "riskScore, riskClass, reason, clipUrl이 저장"],
            ["개인정보", "백엔드 payload/DB 점검", "transcript와 matchedKeywords가 저장되지 않음"],
        ],
        widths=[3.2, 6.7, 6.2],
    )

    doc.add_heading("20. 남은 개선 후보", level=1)
    add_bullets(
        doc,
        [
            "동일 테스트 세트에서 videoOnlyScore 기준 결과와 riskScore 기준 결과를 표로 비교하면 음성 효과를 더 강하게 증명할 수 있다.",
            "클래스별 threshold는 실제 시연 로그를 기반으로 자동 보정하면 오탐과 누락의 균형을 더 잘 맞출 수 있다.",
            "inference_server.py의 대시보드 HTML/JS는 별도 정적 파일로 더 분리하면 유지보수가 쉬워진다.",
            "장시간 1시간 이상 실행 테스트로 WebRTC 연결 누수, CLOVA 사용량, backend queue, FPS 변화를 기록하면 운영 신뢰도가 올라간다.",
            "학습 데이터가 더 확보되면 collapse와 loitering의 hard negative를 추가해 앉아 있음/대기/정상 이동 오탐을 줄일 수 있다.",
        ],
    )

    doc.add_heading("21. 사용자 제공 ZIP PDF 시각화 자료", level=1)
    if user_pdf_visuals:
        add_callout(
            doc,
            "반영 방식",
            "사용자가 제공한 Downloads.zip 안의 PDF 도표를 직접 추출하고 첫 페이지를 PNG로 변환한 뒤, 2장~14장의 본문 대표 도표로 사용했다. 아래 표는 실제 본문에 반영된 원본 PDF 목록이다.",
            LIGHT_BLUE,
        )
        add_table(
            doc,
            ["원본 PDF", "본문 사용 위치", "변환 결과"],
            [
                ["detectwarning_architecture_ppt_white.pdf", "2. 전체 시스템 아키텍처", "PNG 변환 후 그림 1로 삽입"],
                ["detectwarning_sequence_ppt_white.pdf", "4. 실시간 감지 처리 흐름", "PNG 변환 후 그림 2로 삽입"],
                ["detectwarning_inference_pipeline_ppt_white.pdf", "5. 영상 행동 분석 로직", "PNG 변환 후 그림 3으로 삽입"],
                ["detectwarning_risk_flow_ppt_white.pdf", "6. 위험 점수 산정 로직", "PNG 변환 후 그림 4로 삽입"],
                ["detectwarning_audio_flow_ppt_white.pdf", "8. 음성 인식과 개인정보 보호", "PNG 변환 후 그림 5로 삽입"],
                ["detectwarning_cctv_link_sequence_ppt_white.pdf", "10. CCTV 코드와 앱 백엔드 연결", "PNG 변환 후 그림 6으로 삽입"],
                ["detectwarning_webrtc_direct_sequence_ppt_white.pdf", "11. WebRTC direct-first 영상 송출", "PNG 변환 후 그림 7로 삽입"],
                ["detectwarning_clip_flow_ppt_white.pdf", "12. 위험 구간 클립 저장", "PNG 변환 후 그림 8로 삽입"],
                ["detectwarning_er_diagram_ppt_white.pdf", "13. 앱 백엔드 ERD와 데이터 흐름", "PNG 변환 후 그림 9로 삽입"],
                ["detectwarning_backend_event_sequence_ppt_white.pdf", "14. 백엔드 이벤트 저장 및 푸시 알림", "PNG 변환 후 그림 10으로 삽입"],
            ],
            widths=[6.5, 5.2, 4.4],
        )
    else:
        zip_rows: list[list[str]] = []
        if DOWNLOADS_ZIP.exists():
            with zipfile.ZipFile(DOWNLOADS_ZIP, "r") as archive:
                for info in archive.infolist():
                    zip_rows.append([info.filename, f"{info.file_size:,} bytes", "PDF 변환 도구 없음으로 목록만 표시"])
        else:
            zip_rows.append(["Downloads.zip", "없음", "지정된 경로에서 파일을 찾지 못함"])
        add_table(doc, ["파일명", "크기", "비고"], zip_rows, widths=[8.5, 3.0, 5.0])
        add_callout(
            doc,
            "주의",
            "PDF 렌더링 라이브러리를 사용할 수 없어 ZIP 안 PDF를 이미지로 삽입하지 못했다. 이 경우 변환 도구 설치 후 다시 생성해야 한다.",
            LIGHT_RED,
        )

    doc.add_heading("22. 결론", level=1)
    doc.add_paragraph(
        "detectWarning의 최종 구조는 영상 모델 하나에 모든 판단을 맡기는 방식이 아니라, 영상 위험행동, 음성 위험 신호, 시간 안정화, 백엔드 전송 정책, 푸시 알림, 위험 클립 저장을 하나의 실사용 흐름으로 연결한 구조다. "
        "이 설계는 발표에서 두 가지를 명확히 보여준다. 첫째, 영상만으로 판단했을 때의 위험도와 영상+음성을 함께 사용했을 때의 위험도 차이를 수치로 설명할 수 있다. 둘째, 실제 PWA와 백엔드에서 CCTV 코드 연결, 실시간 영상, 위험 기록, 푸시 알림, 클립 재생까지 이어지는 운영 흐름을 시연할 수 있다."
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_report()
