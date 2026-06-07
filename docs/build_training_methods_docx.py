from __future__ import annotations

from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "detectWarning_training_methods_detail.docx"
USER_VISUALS = DOCS / "_user_pdf_visuals_png"
PRESENTATION_VISUALS = DOCS / "presentation_visuals"

FONT = "Malgun Gothic"
INK = "111827"
MUTED = "475569"
BLUE = "2563EB"
BLUE_DARK = "1E3A8A"
GREEN = "16A34A"
RED = "DC2626"
YELLOW = "D97706"
LIGHT_BLUE = "DBEAFE"
LIGHT_GREEN = "DCFCE7"
LIGHT_YELLOW = "FEF3C7"
LIGHT_RED = "FEE2E2"
LIGHT_GRAY = "F1F5F9"


def set_font(run, size: float | None = None, *, bold: bool = False, color: str = INK) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_text(cell, text: str, *, fill: str | None = None, bold: bool = False, color: str = INK, size: float = 8.0) -> None:
    cell.text = ""
    if fill:
        shade(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    set_font(r, size, bold=bold, color=color)


def setup(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(4)

    for style_name, size, color in [
        ("Heading 1", 15.5, BLUE),
        ("Heading 2", 12.0, BLUE_DARK),
        ("Heading 3", 10.5, GREEN),
    ]:
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("detectWarning 학습 기법 상세 문서")
    set_font(r, 8, color=MUTED)


def para(doc: Document, text: str, *, size: float = 9.5, bold: bool = False, color: str = INK) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size, bold=bold, color=color)


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(item)
        set_font(r, 9.0)


def numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(item)
        set_font(r, 9.0)


def table(doc: Document, headers: list[str], rows: list[list[str]], *, widths: list[float] | None = None, font_size: float = 7.6) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell_text(t.cell(0, idx), header, fill=LIGHT_BLUE, bold=True, color=BLUE_DARK, size=8.0)
    for row in rows:
        cells = t.add_row().cells
        for idx, value in enumerate(row):
            cell_text(cells[idx], value, size=font_size)
            if widths and idx < len(widths):
                cells[idx].width = Inches(widths[idx])
    doc.add_paragraph()


def callout(doc: Document, title: str, body: str, *, fill: str = LIGHT_YELLOW) -> None:
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_font(r, 9.5, bold=True, color=BLUE_DARK)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_font(r2, 8.7)
    doc.add_paragraph()


def image(doc: Document, path: Path, caption: str, *, max_width: float = 7.05, max_height: float = 4.2) -> None:
    if not path.exists():
        return
    with Image.open(path) as im:
        w, h = im.size
    aspect = w / max(h, 1)
    width = max_width
    height = width / aspect
    if height > max_height:
        height = max_height
        width = height * aspect
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(path), width=Inches(width), height=Inches(height))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = c.add_run(caption)
    set_font(cr, 7.8, color=MUTED)


def landscape(doc: Document) -> None:
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(11)
    sec.page_height = Inches(8.5)
    sec.top_margin = Inches(0.4)
    sec.bottom_margin = Inches(0.4)
    sec.left_margin = Inches(0.4)
    sec.right_margin = Inches(0.4)


def title(doc: Document) -> None:
    p = doc.add_paragraph()
    r = p.add_run("detectWarning 데이터 학습 기법 상세 설명")
    set_font(r, 25, bold=True)
    p2 = doc.add_paragraph()
    r2 = p2.add_run("AIHub guideline 기반 전처리부터 pose/RGB/I3D/BiLSTM, 자동 튜닝, 앙상블, 오탐 보정까지")
    set_font(r2, 13, bold=True, color=BLUE)
    table(
        doc,
        ["항목", "내용"],
        [
            ["작성일", datetime.now().strftime("%Y-%m-%d")],
            ["문서 목적", "프로젝트에서 데이터를 학습시킬 때 적용한 기법들을 별도로 정리하고, 각 기법을 쓴 이유와 기대 효과를 설명한다."],
            ["학습 목표", "정확도 하나만 올리는 것이 아니라 macro F1, 클래스별 F1, 정상 오탐 감소, 위험 누락 감소, 실시간 추론 안정성을 함께 개선한다."],
            ["최종 추론 클래스", "violence, collapse, loitering. normal은 정상 기준 상태로 쓰며, abduction은 현재 추론 핵심 클래스에서 제외하는 방향이 안정적이다."],
        ],
        widths=[1.5, 5.8],
        font_size=8.2,
    )


def build() -> None:
    doc = Document()
    setup(doc)
    title(doc)

    doc.add_heading("1. 왜 단순 전체 영상 학습이 아니라 clip 기반 학습이 필요한가", level=1)
    para(
        doc,
        "초기에는 전체 영상을 하나의 라벨로 학습시키면 편하지만, AIHub 이상행동 CCTV 데이터는 한 영상 안에 정상 구간과 이상행동 구간이 섞여 있다. "
        "가이드라인상 많은 영상은 정상 1분, 이상행동 2분, 정상 1분처럼 구성되므로 전체 영상을 하나의 이상행동 라벨로 넣으면 정상 프레임이 이상행동으로 학습되는 label noise가 생긴다."
    )
    bullets(
        doc,
        [
            "event 시작/종료 구간을 XML에서 읽어 실제 이상행동 구간 중심으로 학습한다.",
            "앞뒤 10~20초 context를 포함해 폭행 전조, 쓰러지기 직전 자세, 배회 흐름처럼 행동 전후 문맥을 보존한다.",
            "정상 구간은 normal clip으로 따로 만들고, 위험 클래스와 정상 클래스가 같은 manifest 안에서 비교되도록 한다.",
            "한 영상에서 여러 clip을 만들기 때문에 데이터 개수가 늘어나고, 학습 샘플이 실제 판단 단위와 가까워진다.",
        ]
    )
    image(doc, USER_VISUALS / "detectwarning_inference_pipeline_ppt_white.png", "학습/추론 파이프라인 개요: 영상 전체가 아니라 clip 단위로 학습 데이터를 만든다.", max_width=5.5, max_height=7.0)

    doc.add_heading("2. 전체 학습 파이프라인", level=1)
    numbered(
        doc,
        [
            "AIHub filekey를 클래스 분포와 outside/inside 정책에 맞춰 추천하고 다운로드한다.",
            "원본 영상과 XML을 split manifest로 정리한다.",
            "guideline_pose_dataset.py가 XML event 구간과 context를 읽어 guideline clip manifest를 만든다.",
            "pose 추출을 수행하되 person_not_detected가 많아도 RGB/I3D-only fallback으로 샘플을 최대한 살린다.",
            "extract_rgb_video_features.py가 clip 단위 RGB/I3D feature를 추출하고 manifest에 rgb_feature_path, filekey, sample_weight를 기록한다.",
            "action_model.py가 pose sequence 모델을 학습하고, specialized_action_tasks.py가 감지 전용/분류 전용 task를 분리한다.",
            "train_fight_bilstm.py가 fight/noFight 또는 fall/normal 같은 짧은 외부 데이터셋을 보조 booster로 미세조정한다.",
            "auto_tune_action_training.py, ensemble_action_models.py, hybrid_pose_ensemble.py가 여러 trial과 모델 결합을 통해 더 안정적인 설정을 고른다.",
            "학습 결과는 metrics, confusion matrix, class report, threshold, checkpoint 형태로 저장되고 추론 서버가 이를 로드한다.",
        ]
    )

    doc.add_heading("3. 적용한 학습 기법 요약표", level=1)
    table(
        doc,
        ["기법", "해결하려는 문제", "적용 방식", "기대 효과"],
        [
            ["event 구간 자르기 + context", "정상 프레임이 이상행동 라벨로 섞이는 label noise", "XML starttime/duration 기준 event clip 생성, 앞뒤 context 포함", "검증 손실 안정화, 행동 전조 학습, 클래스 특징 선명화"],
            ["clip 단위 샘플 생성", "영상 수가 적고 한 영상이 너무 길어 학습 단위가 둔함", "한 영상에서 여러 위험/정상 clip 생성", "샘플 수 증가, 실시간 추론 단위와 학습 단위 일치"],
            ["normal 클래스 추가", "위험만 학습하면 정상 상황에서 오탐이 커짐", "정상 구간을 normal clip으로 추가", "정상 오탐 감소. 단, normal 비율이 과하면 위험 recall이 낮아질 수 있어 조절 필요"],
            ["pose 실패 샘플 유지", "person_not_detected 때문에 데이터 손실이 큼", "pose_path가 없어도 RGB/I3D-only fallback으로 manifest 유지", "데이터 폐기 감소, 특히 어두운/작은 사람 영상 보존"],
            ["RGB/I3D feature fusion", "pose만으로는 크로마키, 가림, 작은 사람, 카메라 각도에 약함", "clip frame에서 RGB feature를 추출해 pose와 결합", "폭력/쓰러짐처럼 시각 패턴이 중요한 클래스 보완"],
            ["Focal Loss", "쉬운 샘플이 loss를 지배하고 어려운 클래스가 묻힘", "정답 확률이 낮은 어려운 샘플에 더 큰 loss 부여", "소수/난이도 높은 클래스 F1 개선 기대"],
            ["Balanced class weight", "클래스별 샘플 수 불균형", "클래스 빈도 기반 weight 또는 sampler 적용", "macro F1과 소수 클래스 recall 개선"],
            ["Mean/Max pooling", "LSTM/temporal 모델이 불안정할 때 중요한 frame 신호가 묻힘", "시간축 평균과 최대값을 함께 쓰는 pose classifier", "짧은 위험 순간을 놓치지 않고 안정적인 baseline 확보"],
            ["Hard negative mining", "정상인데 위험으로 오탐되는 샘플이 반복됨", "오탐 샘플을 수집해 normal 또는 약한 weight로 재학습 후보화", "false positive 감소"],
            ["Auto tune trial", "수동으로 learning rate, weight, normal ratio 찾기 어려움", "여러 seed/config를 반복 학습하고 best metric 선택", "우연한 seed 성능 편차 감소, 설정 탐색 자동화"],
            ["Ensemble/hybrid", "단일 모델이 특정 클래스에서 흔들림", "pose/RGB/meta/classical model 또는 checkpoint probability 결합", "macro F1 안정화, 클래스별 약점 보완"],
            ["BiLSTM 보조 모델", "짧은 fight/fall 데이터셋은 시간 변화 학습에 적합", "CNN feature + BiLSTM + attention, seed sweep", "violence/collapse 후보를 조건부 booster로 보강"],
        ],
        widths=[1.5, 2.1, 2.5, 2.4],
        font_size=7.3,
    )

    doc.add_heading("4. AIHub guideline 기반 전처리", level=1)
    para(
        doc,
        "가장 큰 성능 저하 원인은 모델 자체보다 데이터 구간이 흐려지는 문제였다. "
        "그래서 학습 데이터는 전체 영상 라벨이 아니라 XML event 구간과 context를 기준으로 다시 구성했다."
    )
    table(
        doc,
        ["항목", "적용 기준", "이유", "담당 파일"],
        [
            ["event clip", "XML event starttime/duration", "이상행동이 실제로 나타나는 구간만 학습한다.", "guideline_pose_dataset.py"],
            ["context", "event 앞뒤 10~20초", "전조/후속 동작을 포함해 행동 흐름을 학습한다.", "guideline_pose_dataset.py"],
            ["normal clip", "event 밖 정상 구간", "정상 상황을 위험과 구분하는 기준을 만든다.", "guideline_pose_dataset.py"],
            ["filekey 추적", "manifest에 filekey/source/xml 기록", "나중에 어떤 filekey가 성능을 망치는지 찾기 쉽다.", "guideline_pose_dataset.py, reporting.py"],
            ["XML matching", "영상명과 XML명 fuzzy/alias 매칭", "파일명 구조 차이로 event 정보가 누락되지 않게 한다.", "guideline_pose_dataset.py"],
        ],
        widths=[1.5, 2.2, 3.0, 2.2],
        font_size=7.7,
    )

    doc.add_heading("5. 클래스 분포와 filekey 선택 정책", level=1)
    para(
        doc,
        "처음에는 모든 클래스를 동일하게 맞추려 했지만, 실제 사용 가능한 filekey 수가 클래스마다 달랐다. "
        "따라서 무리하게 완전 균형을 맞추기보다 abduction을 제외하고 violence/collapse/loitering 데이터를 최대한 확보하되, 학습 때 class weight와 threshold로 보정하는 방향이 더 현실적이었다."
    )
    table(
        doc,
        ["클래스", "정책", "이유"],
        [
            ["violence", "outside 우선, 이후에는 충분히 학습 가능하도록 포함", "폭력은 실외 CCTV 맥락이 중요하고 외부 fight 데이터셋으로 추가 보강 가능하다."],
            ["collapse", "outside + inside 모두 활용 가능", "쓰러짐은 장소보다 자세 변화가 더 중요하므로 데이터 확보가 우선이다."],
            ["loitering", "outside + inside 모두 활용 가능", "배회는 장면 맥락 영향을 받지만 데이터 수가 부족하면 실내도 보조로 쓸 수 있다."],
            ["abduction", "현재 추론 핵심 클래스에서 제외", "데이터 수와 구분성이 부족하면 전체 macro F1과 혼동 행렬을 악화시킬 수 있다."],
            ["normal", "25% 안팎에서 시작, 오탐/누락에 따라 조정", "normal이 너무 적으면 오탐, 너무 많으면 위험 recall 저하가 생긴다."],
        ],
        widths=[1.3, 2.9, 4.7],
        font_size=7.8,
    )

    doc.add_heading("6. pose 기반 학습 기법", level=1)
    para(
        doc,
        "pose는 관절점 시퀀스를 사용하므로 배경 변화에 상대적으로 강하고, 실시간 계산 비용이 RGB 비디오 모델보다 낮다. "
        "다만 사람 검출 실패, 작은 사람, 가림, 앉은 자세/쓰러짐 혼동에 취약하므로 fallback과 보조 RGB/I3D가 필요했다."
    )
    table(
        doc,
        ["구성", "설명", "정확도/F1에 기대한 영향"],
        [
            ["PoseSequenceDataset", "pose_path가 있는 npz를 읽어 sequence tensor와 mask, label을 만든다.", "관절점 기반 baseline을 안정적으로 학습한다."],
            ["TemporalPoseClassifier", "시간축 pose 변화를 모델링하는 기본 classifier다.", "행동 흐름이 있는 violence/collapse/loitering 구분에 사용한다."],
            ["MeanMaxTemporalPoseClassifier", "시간 평균과 최대 pooling을 함께 사용한다.", "짧은 위험 순간과 전체 자세 흐름을 함께 반영한다."],
            ["FocalLoss", "어려운 샘플과 소수 클래스의 loss 비중을 키운다.", "macro F1과 소수 클래스 recall 개선을 노린다."],
            ["Balanced weight/sampler", "클래스 빈도에 따라 loss나 샘플링을 보정한다.", "데이터 불균형으로 특정 클래스만 맞추는 문제를 줄인다."],
            ["RGB-only fallback weight", "pose 실패지만 RGB feature가 있는 샘플은 낮은 신뢰도 가중치로 살린다.", "데이터 폐기를 줄이되 품질 낮은 샘플이 학습을 지배하지 않게 한다."],
        ],
        widths=[2.0, 4.0, 3.0],
        font_size=7.7,
    )

    doc.add_heading("7. RGB/I3D feature fusion", level=1)
    para(
        doc,
        "RGB/I3D를 도입한 이유는 pose만으로는 사람 검출 실패나 작은 동작, 카메라 각도, 배경 맥락을 충분히 보지 못하기 때문이다. "
        "반대로 RGB만 쓰면 배경/조명/크로마키에 민감해지므로 pose와 fusion하는 구조가 더 안전하다."
    )
    table(
        doc,
        ["항목", "설명", "주의점"],
        [
            ["clip별 RGB feature", "각 guideline clip에서 일정 frame을 샘플링해 feature npz를 만든다.", "source video가 없으면 missing source로 기록하고 추적해야 한다."],
            ["I3D/R3D 모델 분리", "실행별/model별 feature 경로를 분리해 서로 다른 feature가 섞이지 않게 한다.", "i3d_r50 실패 시 fallback이 기록되어야 한다."],
            ["filekey tracking", "feature manifest에 filekey와 source를 기록한다.", "어떤 filekey가 성능을 낮추는지 분석 가능하다."],
            ["pose와 결합", "pose score, RGB feature, bbox/meta를 hybrid model 또는 추론 score에 반영한다.", "두 modality label set이 다르면 후처리에서 skip될 수 있다."],
            ["크로마키", "단독 학습보다 augmentation 성격으로만 사용한다.", "배경이 비현실적이면 실제 CCTV 성능을 떨어뜨릴 수 있다."],
        ],
        widths=[1.7, 4.2, 3.0],
        font_size=7.7,
    )

    doc.add_heading("8. 감지 전용 학습과 분류 전용 학습 분리", level=1)
    para(
        doc,
        "프로젝트의 실제 목적은 이상행동 세부 분류만이 아니라 위험상황 감지다. 그래서 normal 포함 감지 task와 abnormal끼리의 세부 분류 task를 분리하는 구조가 필요했다."
    )
    table(
        doc,
        ["task", "라벨 구성", "목적", "이유"],
        [
            ["감지 전용", "normal vs abnormal 또는 normal 포함 4클래스", "정상 오탐을 줄이고 위험 후보를 찾는다.", "실사용에서는 정상 상황을 위험으로 띄우지 않는 것이 중요하다."],
            ["분류 전용", "violence/collapse/loitering", "위험이라고 판단된 뒤 어떤 위험인지 구분한다.", "normal 축이 혼동 행렬을 지배하면 위험 클래스별 F1 해석이 흐려진다."],
            ["추론 결합", "감지 score + class score", "위험 여부와 위험 종류를 따로 보고 합친다.", "오탐과 누락을 동시에 줄이기 쉽다."],
        ],
        widths=[1.5, 2.3, 3.0, 3.0],
        font_size=7.8,
    )

    doc.add_heading("9. BiLSTM + Attention 미세조정", level=1)
    para(
        doc,
        "외부의 짧은 fight/noFight 또는 fall/normal 데이터셋은 2초 내외 영상에서 frame feature의 시간 변화를 학습하기 좋다. "
        "이때 CNN은 각 프레임 특징을 뽑고, BiLSTM은 앞뒤 시간 관계를 학습하며, attention은 중요한 프레임에 더 큰 가중치를 준다."
    )
    table(
        doc,
        ["구성", "역할", "프로젝트 적용 방식"],
        [
            ["CNN feature extractor", "각 frame의 공간적 특징 추출", "작은 데이터셋에서도 빠르게 frame-level feature를 만든다."],
            ["BiLSTM", "이전/이후 frame 관계를 양방향으로 학습", "싸움 동작처럼 움직임 변화가 핵심인 상황에 적합하다."],
            ["Self-attention", "중요 frame에 더 집중", "모든 frame이 같은 중요도를 갖지 않으므로 결정적인 동작을 강조한다."],
            ["Seed sweep", "여러 seed 학습 후 best checkpoint 선택", "작은 데이터셋의 우연한 split/seed 성능 편차를 줄인다."],
            ["보조 점수화", "주 모델 대체가 아니라 조건부 booster", "기존 3클래스 모델의 점수를 비정상적으로 키우지 않게 한다."],
        ],
        widths=[1.7, 3.0, 4.2],
        font_size=7.8,
    )
    callout(
        doc,
        "왜 보조 모델로만 쓰는가",
        "fight/noFight 데이터셋은 violence에는 강하지만 collapse/loitering을 대표하지 않는다. 따라서 기존 3클래스 모델을 덮어쓰기보다, violence 후보가 이미 있을 때 추가 확인 신호로 쓰는 것이 오탐을 줄이는 데 더 안전하다.",
        fill=LIGHT_GREEN,
    )

    doc.add_heading("10. 자동 튜닝, 앙상블, hybrid", level=1)
    table(
        doc,
        ["기법", "동작", "효과", "주의점"],
        [
            ["Auto-tune trial", "learning rate, class weight, focal 설정, seed, normal ratio를 바꿔 반복 학습", "수동 설정 의존도를 줄이고 best validation metric을 찾는다.", "trial 수가 늘면 시간이 길어진다."],
            ["Seed sweep", "같은 설정을 여러 seed로 학습", "작은 데이터셋에서 우연한 결과를 줄인다.", "검증 set이 작으면 여전히 과대평가될 수 있다."],
            ["Checkpoint ensemble", "여러 모델 probability를 평균/가중 평균", "단일 모델의 클래스별 흔들림을 줄인다.", "라벨 구성이 다르면 결합하면 안 된다."],
            ["Hybrid model", "pose/RGB/meta feature matrix로 extra trees 등 classical model 탐색", "neural model이 놓친 meta 패턴을 보완한다.", "RGB-only row나 label conflict row 필터링이 필요하다."],
            ["Threshold calibration", "클래스별 전송 threshold와 posterior threshold 조정", "정상 오탐과 위험 누락의 균형을 맞춘다.", "학습 정확도보다 실사용 false positive를 우선 봐야 한다."],
        ],
        widths=[1.7, 3.0, 2.4, 2.2],
        font_size=7.7,
    )

    doc.add_heading("11. 과적합과 검증 손실이 높을 때의 해석", level=1)
    para(
        doc,
        "학습 손실은 줄어드는데 검증 손실이 그대로거나 높으면 모델이 학습 데이터의 세부 패턴만 외우고 검증 데이터 일반화에 실패한 것이다. "
        "이 프로젝트에서는 label noise, class imbalance, normal 비율, pose 실패 샘플, RGB feature 누락, filekey별 domain 차이가 주요 원인이 될 수 있다."
    )
    table(
        doc,
        ["증상", "가능 원인", "대응"],
        [
            ["train loss 낮고 val loss 높음", "과적합, split domain 차이, label noise", "weight decay, dropout, early stopping, hard negative, XML 매칭 점검"],
            ["특정 클래스 F1 낮음", "클래스 데이터 부족, threshold 과함, label 혼동", "class weight, focal loss, 해당 클래스 데이터 추가, class별 threshold 조정"],
            ["normal 오탐 많음", "normal 비율 낮음, hard negative 부족", "normal clip 비율 25~35% 조정, hard negative mining"],
            ["위험 recall 낮음", "normal이 너무 많거나 threshold가 높음", "normal 비율 낮춤, danger threshold 완화, audio/video 결합 강화"],
            ["pose accuracy 낮음", "person_not_detected, 작은 사람, 가림", "RGB/I3D fallback, detector threshold/imgsz 재시도, 낮은 pose 품질 weight 조정"],
        ],
        widths=[2.0, 3.1, 3.8],
        font_size=7.7,
    )

    doc.add_heading("12. 최종적으로 추천하는 학습 운영 방식", level=1)
    bullets(
        doc,
        [
            "1차: guideline XML 기반 clip + normal + pose/RGB feature를 최대한 많이 추출한다. 이때 원본 삭제보다 manifest와 XML 추적성이 더 중요하다.",
            "2차: pose baseline을 focal loss, balanced weight, mean/max pooling으로 학습해 3클래스 기준 성능을 본다.",
            "3차: RGB/I3D feature를 결합해 pose가 약한 clip을 보완한다.",
            "4차: normal 포함 감지 task와 violence/collapse/loitering 분류 task를 분리해 평가한다.",
            "5차: fight/fall 같은 작은 외부 데이터셋은 기존 모델을 덮어쓰지 말고 조건부 booster로만 쓴다.",
            "6차: auto-tune과 seed sweep으로 best checkpoint를 고르고, threshold calibration으로 실제 오탐을 줄인다.",
            "7차: 실제 시연 영상에서 hard negative를 수집해 normal/weak negative로 재학습하면 발표용 안정성이 가장 크게 좋아진다.",
        ]
    )
    image(doc, PRESENTATION_VISUALS / "04_class_guardrails.png", "클래스별 오탐 방지 기준: 학습 성능뿐 아니라 실사용 threshold가 중요하다.", max_width=7.05, max_height=4.0)

    landscape(doc)
    doc.add_heading("13. 학습 기법과 담당 코드 매핑", level=1)
    table(
        doc,
        ["기법/기능", "담당 파일", "입력", "출력", "실사용 영향"],
        [
            ["AIHub 다운로드/병합", "action_training_pipeline.py, predownload_aihub_filekey.py", "datasetkey/filekey", "raw manifest", "학습 데이터 확보 속도"],
            ["filekey 추천", "dashboard_aihub.py", "AIHub filekey 목록, class 정책", "다음 작업 filekey", "클래스 분포와 outside/inside 정책 반영"],
            ["guideline clip", "guideline_pose_dataset.py", "raw/split manifest, XML", "guideline_prepared_*.jsonl", "label noise 감소"],
            ["pose prepare", "action_training_pipeline.py, pipeline_prepare.py", "clip/source video", "pose npz, pose_path", "pose baseline 품질"],
            ["RGB/I3D feature", "extract_rgb_video_features.py", "guideline clip, source video", "feature npz, rgb_feature_path", "pose 실패/배경 맥락 보완"],
            ["pose classifier", "action_model.py", "prepared manifest", "best_action_model.pt, metrics.json", "3클래스 행동 분류"],
            ["전용 task", "specialized_action_tasks.py", "normal/abnormal manifest", "detection/classification artifacts", "감지와 분류 목적 분리"],
            ["BiLSTM fine-tune", "train_fight_bilstm.py", "fight/fall external dataset", "booster checkpoint", "violence/collapse 보조 확인 신호"],
            ["seed sweep", "fight_bilstm_seed_sweep.py", "seed config set", "best seed checkpoint", "작은 데이터셋 안정화"],
            ["auto tune", "auto_tune_action_training.py", "trial config", "best config/artifacts", "하이퍼파라미터 탐색"],
            ["ensemble", "ensemble_action_models.py", "여러 checkpoint", "ensemble summary/model", "macro F1 안정화"],
            ["hybrid", "hybrid_pose_ensemble.py", "pose/RGB/meta feature", "hybrid model summary", "서로 다른 feature 결합"],
            ["결과 진단", "training_insights.py, reporting.py", "metrics, confusion matrix", "insight/report", "다음 개선 방향 결정"],
        ],
        widths=[1.7, 2.4, 2.0, 2.2, 2.4],
        font_size=6.8,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
